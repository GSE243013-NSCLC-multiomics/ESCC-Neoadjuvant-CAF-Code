#!/usr/bin/env python3

from pathlib import Path
from docx import Document
import hashlib
import shutil

ROOT = Path.cwd()

INPUT = (
    ROOT /
    "04_results/manuscript_update/Step23/"
    "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1C_QA2E_LOCKED_EDITS.docx"
)

OUTPUT = (
    ROOT /
    "04_results/manuscript_update/Step23/"
    "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1F_VERIFIED_REFERENCES.docx"
)

AUDIT = (
    ROOT /
    "04_results/manuscript_update/Step23/"
    "STEP23C_QA2B1F_REFERENCE_INSERT_AUDIT.txt"
)

CHECKSUM = (
    ROOT /
    "00_metadata/Step23_manuscript_update/"
    "STEP23C_QA2B1F_REFERENCE_INSERT_CHECKSUMS.sha256"
)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


if not INPUT.exists():
    raise SystemExit(f"STOP: input missing:\n{INPUT}")

input_sha_before = sha256(INPUT)

print("=" * 88)
print("STEP23C-QA2B1F: INSERT VERIFIED DATASET-SOURCE REFERENCES")
print("=" * 88)
print()
print("INPUT:")
print(" ", INPUT.relative_to(ROOT))
print()
print("INPUT SHA256:")
print(" ", input_sha_before)
print()


# ------------------------------------------------------------------
# 1. Make a NEW copy only
# ------------------------------------------------------------------

shutil.copy2(INPUT, OUTPUT)

doc = Document(OUTPUT)

paras = doc.paragraphs

ref_hits = [
    i for i, p in enumerate(paras)
    if p.text.strip() == "References"
]

fig_hits = [
    i for i, p in enumerate(paras)
    if p.text.strip() == "Figure Legends"
]

if len(ref_hits) != 1:
    raise SystemExit(
        f"STOP: expected exactly one References heading; found {len(ref_hits)}"
    )

if len(fig_hits) != 1:
    raise SystemExit(
        f"STOP: expected exactly one Figure Legends heading; found {len(fig_hits)}"
    )

ref_i = ref_hits[0]
fig_i = fig_hits[0]

if not ref_i < fig_i:
    raise SystemExit(
        "STOP: References heading does not precede Figure Legends."
    )

between = doc.paragraphs[ref_i + 1:fig_i]

print("CURRENT REFERENCES BLOCK:")
for p in between:
    print(" ", repr(p.text))
print()


# ------------------------------------------------------------------
# 2. Guard against accidental prior insertion
# ------------------------------------------------------------------

whole_text = "\n".join(p.text for p in doc.paragraphs)

dois = [
    "10.1016/j.isci.2023.106480",
    "10.1038/s41392-023-01518-0",
    "10.1186/s13073-024-01320-9",
]

for doi in dois:
    if doi in whole_text:
        raise SystemExit(
            f"STOP: DOI already present before QA2B1F insertion: {doi}"
        )


# ------------------------------------------------------------------
# 3. Verify placeholder block
# ------------------------------------------------------------------

placeholder_candidates = [
    p for p in between
    if (
        "References require human completion" in p.text
        or "[CITATION NEEDED]" in p.text
        or "placeholders apply where literature support is required" in p.text
    )
]

if len(placeholder_candidates) != 1:
    print("STOP: reference placeholder block was not uniquely identified.")
    print("Candidates found:", len(placeholder_candidates))
    raise SystemExit(1)

placeholder = placeholder_candidates[0]
base_style = placeholder.style


# ------------------------------------------------------------------
# 4. Verified working-copy reference strings
#
# Neutral biomedical format.
# Final target-journal formatting remains a later packaging step.
# ------------------------------------------------------------------

refs = [
    (
        "Okuda S, Ohuchida K, Nakamura S, et al. "
        "Neoadjuvant chemotherapy enhances anti-tumor immune response "
        "of tumor microenvironment in human esophageal squamous cell carcinoma. "
        "iScience. 2023;26(4):106480. "
        "doi:10.1016/j.isci.2023.106480."
    ),
    (
        "Yang Y, Li Y, Yu H, et al. "
        "Comprehensive landscape of resistance mechanisms for neoadjuvant "
        "therapy in esophageal squamous cell carcinoma by single-cell transcriptomics. "
        "Signal Transduct Target Ther. 2023;8(1):298. "
        "doi:10.1038/s41392-023-01518-0."
    ),
    (
        "Ji G, Yang Q, Wang S, et al. "
        "Single-cell profiling of response to neoadjuvant chemo-immunotherapy "
        "in surgically resectable esophageal squamous cell carcinoma. "
        "Genome Med. 2024;16(1):49. "
        "doi:10.1186/s13073-024-01320-9."
    ),
]

incomplete_note = (
    "[Additional manuscript-wide literature references remain to be "
    "completed and verified before final submission.]"
)


# ------------------------------------------------------------------
# 5. Replace only the existing placeholder paragraph,
#    then insert remaining entries immediately before Figure Legends.
# ------------------------------------------------------------------

placeholder.text = refs[0]
placeholder.style = base_style

# Re-fetch Figure Legends paragraph because paragraph list may shift
figure_heading = next(
    p for p in doc.paragraphs
    if p.text.strip() == "Figure Legends"
)

for txt in refs[1:]:
    p = figure_heading.insert_paragraph_before(txt)
    p.style = base_style

note_p = figure_heading.insert_paragraph_before(incomplete_note)
note_p.style = base_style


# ------------------------------------------------------------------
# 6. Save
# ------------------------------------------------------------------

doc.save(OUTPUT)


# ------------------------------------------------------------------
# 7. Structural regression QA
# ------------------------------------------------------------------

check = Document(OUTPUT)
text = "\n".join(p.text for p in check.paragraphs)

failures = []

for doi in dois:
    n = text.count(doi)
    print(f"DOI COUNT {doi}: {n}")
    if n != 1:
        failures.append(
            f"{doi}: expected 1 occurrence, found {n}"
        )

for phrase in [
    "Okuda S, Ohuchida K, Nakamura S, et al.",
    "Yang Y, Li Y, Yu H, et al.",
    "Ji G, Yang Q, Wang S, et al.",
    "Additional manuscript-wide literature references remain",
]:
    if phrase not in text:
        failures.append(
            f"missing required reference text: {phrase}"
        )

legacy_placeholder = (
    "References require human completion; "
    "no fabricated citations are included"
)

if legacy_placeholder in text:
    failures.append(
        "old References placeholder remains present"
    )

# Do not permit accidental deletion of downstream section.
if "Figure Legends" not in text:
    failures.append(
        "Figure Legends heading missing after edit"
    )

if failures:
    print()
    print("STOP: POST-SAVE QA FAILED")
    for x in failures:
        print(" -", x)

    OUTPUT.unlink(missing_ok=True)
    raise SystemExit(1)


# ------------------------------------------------------------------
# 8. Verify source remained frozen
# ------------------------------------------------------------------

input_sha_after = sha256(INPUT)

if input_sha_after != input_sha_before:
    OUTPUT.unlink(missing_ok=True)
    raise SystemExit(
        "CRITICAL STOP: QA2B1C source DOCX changed."
    )


# ------------------------------------------------------------------
# 9. Audit
# ------------------------------------------------------------------

output_sha = sha256(OUTPUT)

audit_lines = [
    "STEP23C-QA2B1F VERIFIED REFERENCE INSERT AUDIT",
    "=" * 80,
    "",
    f"INPUT: {INPUT.relative_to(ROOT)}",
    f"INPUT SHA256: {input_sha_before}",
    "",
    f"OUTPUT: {OUTPUT.relative_to(ROOT)}",
    f"OUTPUT SHA256: {output_sha}",
    "",
    "SOURCE DOCX MODIFIED: NO",
    "STEP22A RECOMPUTED: NO",
    "GLOBAL FIND/REPLACE: NO",
    "",
    "QA2C-054 DATASET-SOURCE REFERENCE HOLD:",
    "RESOLVED FOR GSE197677 / GSE221561 / OMIX005710",
    "",
    "INSERTED VERIFIED SOURCE PUBLICATIONS:",
    "1. GSE197677 — Okuda et al., iScience 2023",
    "   DOI 10.1016/j.isci.2023.106480",
    "2. GSE221561 — Yang et al., Signal Transduct Target Ther 2023",
    "   DOI 10.1038/s41392-023-01518-0",
    "3. OMIX005710 — Ji et al., Genome Med 2024",
    "   DOI 10.1186/s13073-024-01320-9",
    "",
    "IMPORTANT:",
    "The manuscript-wide literature bibliography is NOT declared complete.",
    "An explicit remaining-reference note is retained in the working copy.",
    "",
    "STATUS:",
    "STEP23C_QA2B1F_VERIFIED_SOURCE_REFERENCES_INSERTED",
    "",
]

AUDIT.write_text(
    "\n".join(audit_lines),
    encoding="utf-8"
)

CHECKSUM.parent.mkdir(
    parents=True,
    exist_ok=True
)

with open(CHECKSUM, "w", encoding="utf-8") as f:
    for p in [INPUT, OUTPUT, AUDIT]:
        f.write(
            f"{sha256(p)}  {p.relative_to(ROOT)}\n"
        )


# ------------------------------------------------------------------
# 10. Console summary
# ------------------------------------------------------------------

print()
print("=" * 88)
print("STEP23C-QA2B1F COMPLETE")
print("=" * 88)
print()
print("QA2B1C SOURCE MODIFIED          : NO")
print("STEP22A RECOMPUTED              : NO")
print("GLOBAL FIND/REPLACE             : NO")
print()
print("VERIFIED REFERENCES INSERTED    : 3 / 3")
print("GSE197677                       : INSERTED")
print("GSE221561                       : INSERTED")
print("OMIX005710                      : INSERTED")
print()
print("QA2C-054 SOURCE HOLD            : RESOLVED")
print("FULL BIBLIOGRAPHY COMPLETE      : NO")
print()
print("OUTPUT DOCX:")
print(" ", OUTPUT.relative_to(ROOT))
print()
print("AUDIT:")
print(" ", AUDIT.relative_to(ROOT))
print()
print("STATUS:")
print("STEP23C_QA2B1F_VERIFIED_SOURCE_REFERENCES_INSERTED")
print()
print("NEXT:")
print(
    "RENDER QA2B1F DOCX AND PERFORM FINAL REFERENCES + PAGE-LAYOUT QA."
)
print()

