#!/usr/bin/env python3

from pathlib import Path
from docx import Document
import hashlib

ROOT = Path.cwd()

DOCX = (
    ROOT /
    "04_results/manuscript_update/Step23/"
    "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1B_CONTEXTUAL_EDITS.docx"
)

if not DOCX.exists():
    raise SystemExit(f"STOP: missing input DOCX:\n{DOCX}")

doc = Document(DOCX)

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
    return h.hexdigest()

def norm(s):
    return " ".join(
        s.replace("\u2013", "-")
         .replace("\u2014", "-")
         .split()
    ).lower()

def find_all(fragment):
    q = norm(fragment)
    out = []
    for i, p in enumerate(doc.paragraphs, start=1):
        if q in norm(p.text):
            out.append(i)
    return out

def show(idx, radius=1):
    lo = max(1, idx - radius)
    hi = min(len(doc.paragraphs), idx + radius)

    for j in range(lo, hi + 1):
        print(f"[{j}]")
        print(repr(doc.paragraphs[j - 1].text))
        print()

checks = [
    (
        "QA2C-004",
        "limited by cohort-specific noise. We asked whether a reproducible"
    ),
    (
        "QA2C-006",
        "A conserved fibroblast CORE2 response was identified across cohorts"
    ),
    (
        "QA2C-009",
        "These results nominate a conserved fibroblast CORE2 response as a"
    ),
    (
        "QA2C-010",
        "reproducible treatment-associated stromal feature in ESCC"
    ),
    (
        "QA2C-011",
        "In an independent neoadjuvant single-cell cohort (OMIX005710)"
    ),
    (
        "QA2C-029",
        "Definitive ESCC-only external validation in OMIX005710"
    ),
    (
        "QA2C-040",
        "The OMIX005710 analysis provides an important negative external finding"
    ),
    (
        "QA2C-042",
        "A key limitation of the OMIX005710 analysis"
    ),
    (
        "QA2C-048",
        "14. External validation in OMIX005710"
    ),
    (
        "QA2C-049",
        "External validation was additionally assessed using the publicly available OMIX005710"
    ),
    (
        "QA2C-050",
        "The source cohort contained one patient with esophageal adenosquamous carcinoma"
    ),
    (
        "QA2C-051",
        "Supplementary Figure S2: Discordant / non-conserved signaling candidates"
    ),
]

print()
print("=" * 100)
print("STEP23C-QA2B1C0: QA2E LOCKED-DECISION PREFLIGHT")
print("=" * 100)
print()

print("INPUT:")
print(" ", DOCX.relative_to(ROOT))
print()

print("SHA256:")
print(" ", sha256(DOCX))
print()

print("DOCX PARAGRAPHS:", len(doc.paragraphs))
print()

problems = []

for item_id, fragment in checks:
    hits = find_all(fragment)

    print("=" * 100)
    print(item_id)
    print("=" * 100)
    print("SEARCH FRAGMENT:")
    print(fragment)
    print()
    print("MATCHES:", hits)
    print()

    if len(hits) == 1:
        show(hits[0], radius=1)
    else:
        problems.append((item_id, hits))

# ------------------------------------------------------------
# Data Availability
# ------------------------------------------------------------

print("=" * 100)
print("QA2C-052 — DATA AVAILABILITY")
print("=" * 100)

da_hits = []

for i, p in enumerate(doc.paragraphs, start=1):
    t = norm(p.text)

    if t in {
        "data availability",
        "data availability:"
    }:
        da_hits.append(i)

print("HEADING MATCHES:", da_hits)
print()

for i in da_hits:
    show(i, radius=2)

if len(da_hits) != 1:
    problems.append(("QA2C-052_DATA_AVAILABILITY", da_hits))


# ------------------------------------------------------------
# Supplementary Material Index
# ------------------------------------------------------------

print("=" * 100)
print("QA2C-053 — SUPPLEMENTARY MATERIAL INDEX")
print("=" * 100)

sm_hits = []

for i, p in enumerate(doc.paragraphs, start=1):
    t = norm(p.text)

    if t in {
        "supplementary material index",
        "supplementary materials index",
        "supplementary material"
    }:
        sm_hits.append(i)

print("HEADING MATCHES:", sm_hits)
print()

for i in sm_hits:
    show(i, radius=7)

if len(sm_hits) != 1:
    problems.append(("QA2C-053_SUPPLEMENTARY_INDEX", sm_hits))


# ------------------------------------------------------------
# References hold
# ------------------------------------------------------------

print("=" * 100)
print("QA2C-054 — REFERENCES")
print("=" * 100)

ref_hits = []

for i, p in enumerate(doc.paragraphs, start=1):
    if norm(p.text) == "references":
        ref_hits.append(i)

print("HEADING MATCHES:", ref_hits)
print()

for i in ref_hits:
    show(i, radius=2)

print("LOCKED ACTION:")
print("HOLD_FOR_BIBLIOGRAPHIC_VERIFICATION")
print("NO CITATION TEXT WILL BE GENERATED OR INSERTED IN QA2B1C.")
print()


print("=" * 100)
print("PREFLIGHT SUMMARY")
print("=" * 100)
print()

print("QA1B / STEP22A MODIFIED       : NO")
print("CURRENT DOCX MODIFIED         : NO")
print("AUTOMATIC REWRITES            : NO")
print("GLOBAL FIND/REPLACE           : NO")
print()

if problems:
    print("STATUS:")
    print("STEP23C_QA2B1C0_PREFLIGHT_NEEDS_MAPPING_REVIEW")
    print()
    print("REVIEW ITEMS:")
    for item, hits in problems:
        print(f"  {item}: matches={hits}")
else:
    print("STATUS:")
    print("STEP23C_QA2B1C0_PREFLIGHT_PASS")

print()
print("NEXT:")
print(
    "IF PASS, APPLY QA2E 01-14 APPROVED TEXT TO A NEW COPY; "
    "KEEP QA2C-054 REFERENCES ON HOLD."
)
print()
