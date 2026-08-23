#!/usr/bin/env python3

from pathlib import Path
from copy import deepcopy
import csv
import hashlib
import shutil

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


ROOT = Path.cwd()

RESULT = ROOT / "04_results" / "manuscript_update" / "Step23"
META   = ROOT / "00_metadata" / "Step23_manuscript_update"

INPUT = (
    RESULT /
    "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1B_CONTEXTUAL_EDITS.docx"
)

OUTPUT = (
    RESULT /
    "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1C_QA2E_LOCKED_EDITS.docx"
)

OUT_AUDIT = META / "STEP23C_QA2B1C_QA2E_EDIT_AUDIT.csv"
OUT_SHA   = META / "STEP23C_QA2B1C_CHECKSUMS.sha256"
OUT_HOLD  = RESULT / "STEP23C_QA2B1C_REFERENCE_VERIFICATION_HOLD.txt"


# ======================================================================
# Approved QA2E text — copied from frozen human decision dossier
# ======================================================================

A004 = (
    "We asked whether a treatment-associated fibroblast response could be "
    "reproducibly identified across two discovery ESCC single-cell RNA-seq "
    "cohorts at the sample/patient level."
)

A006 = (
    "Within the two discovery cohorts, a fibroblast CORE2 response showed "
    "reproducible same-direction treatment-associated effects."
)

A009_010 = (
    "These results nominate a fibroblast CORE2 response that was reproducible "
    "across the two discovery cohorts, while the independent OMIX005710 "
    "secondary sensitivity analysis showed directional discordance; upstream "
    "regulator and extracellular signaling candidates remain hypothesis-generating."
)

A011 = (
    "In an independent neoadjuvant single-cell cohort (OMIX005710), definitive "
    "primary ESCC-only validation was precluded because the single patient with "
    "esophageal adenosquamous carcinoma could not be identified from the "
    "available patient-level metadata. In a pre-specified secondary "
    "histology-sensitivity analysis spanning all 22 possible EASC assignments, "
    "the treatment-associated CORE direction was consistently opposite to that "
    "observed in the discovery cohorts."
)

A029 = (
    "Definitive primary ESCC-only external validation in OMIX005710 was "
    "unavailable because the single EASC case could not be mapped to an "
    "individual patient using the available public metadata. We therefore "
    "evaluated the pre-specified secondary histology-sensitivity analysis across "
    "all 22 possible EASC assignments, without selecting or promoting any "
    "individual scenario."
)

A040 = (
    "The OMIX005710 analysis provides an important secondary external finding. "
    "Although unresolved patient-level histology precluded definitive primary "
    "ESCC-only validation, the pre-specified secondary histology-sensitivity "
    "analysis consistently yielded a treatment-associated CORE direction "
    "opposite to that observed in the two discovery cohorts. This directional "
    "discordance was robust to all hypothetical EASC assignments, independent "
    "reconstruction of the scoring pipeline, and paired leave-one-patient-out "
    "analyses, arguing against a simple sign, pairing, reference-set, or "
    "single-patient artifact."
)

A042 = (
    "A key limitation of the OMIX005710 analysis was the inability to identify "
    "the single EASC patient at the individual-patient level from the available "
    "public metadata, which precluded construction of a definitive primary "
    "ESCC-only validation set. We therefore relied on a pre-specified secondary "
    "sensitivity analysis across all 22 possible EASC assignments. In addition, "
    "only five paired patients satisfied the frozen fibroblast technical criteria "
    "before hypothetical EASC exclusion, and scenarios in which a paired patient "
    "was assigned as EASC contained four evaluable pairs. This small paired sample "
    "size limits formal inferential resolution; accordingly, interpretation "
    "focused on direction, patient-level consistency, scenario robustness, and "
    "leave-one-patient-out influence rather than statistical significance."
)

A048 = (
    "14. OMIX005710 external-validation framework and secondary "
    "histology-sensitivity analysis"
)

A049 = (
    "A third independent neoadjuvant single-cell RNA-sequencing cohort, "
    "OMIX005710, was evaluated under a pre-specified external-validation "
    "framework. Fibroblast pseudobulk profiles were constructed from raw counts "
    "after broad cell-type annotation using a marker framework independent of "
    "the CORE genes. A pre-specified minimum of 20 fibroblast cells per "
    "patient-timepoint was required for technical eligibility; one sample "
    "(P15 before treatment, 18 fibroblasts) therefore remained in raw provenance "
    "but was excluded from the technical reference and paired analysis. Raw "
    "pseudobulk counts were transformed to library-size CPM and log2(CPM+1), "
    "followed by gene-wise z-scoring within the eligible OMIX tumor reference. "
    "CORE2_STRONG and CORE4 were calculated using the previously frozen formulas "
    "without refitting gene weights."
)

A050 = (
    "The source cohort contained one patient with esophageal adenosquamous "
    "carcinoma (EASC), but the corresponding individual patient could not be "
    "resolved from the publicly available metadata. Accordingly, definitive "
    "primary ESCC-only validation was classified as unavailable. Before reading "
    "any CORE scores, we therefore pre-specified a secondary "
    "histology-sensitivity framework comprising all 22 possible assignments of "
    "the EASC identity. Each scenario excluded the corresponding patient's "
    "technically eligible tumor observations from the z-score reference and, "
    "when applicable, from the paired analysis. Paired treatment effects were "
    "defined as after-treatment minus before-treatment score. Robustness to "
    "individual paired patients was assessed by full leave-one-patient-out "
    "recomputation, with the omitted patient removed from both the z-score "
    "reference and the paired effect calculation."
)

A051 = (
    "Supplementary Figure S2: Discordant signaling candidates across the two "
    "discovery cohorts"
)

A052 = (
    "GSE197677 and GSE221561 are publicly available single-cell RNA-seq datasets. "
    "OMIX005710, used for the secondary external histology-sensitivity analysis, "
    "is publicly available under accession OMIX005710. No new sequencing data "
    "were generated in this study. Accession and repository details will be "
    "provided in the final submission."
)

A053 = [
    (
        "Supplementary Table S13: OMIX005710 sample metadata, technical "
        "eligibility, and patient/timepoint composition."
    ),
    (
        "Supplementary Table S14: Pre-specified 22-scenario histology-sensitivity "
        "results for CORE2_STRONG, CORE4, and component directions."
    ),
    (
        "Supplementary Table S15: Paired leave-one-patient-out robustness "
        "analysis and independent score-recalculation audit."
    ),
    (
        "Supplementary Note S1: OMIX005710 histology-resolution limitation, "
        "primary-validation gate, and analysis governance."
    ),
]


# ======================================================================
# Utilities
# ======================================================================

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def norm(s):
    return " ".join(
        s.replace("\u2013", "-")
         .replace("\u2014", "-")
         .replace("\u00a0", " ")
         .split()
    ).lower()


def set_paragraph_text(p, text):
    """
    Preserve paragraph style/properties and first-run formatting.
    Remove textual contents of remaining runs.
    """
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def delete_paragraph(p):
    el = p._element
    parent = el.getparent()
    if parent is None:
        raise RuntimeError("PARAGRAPH_ALREADY_DETACHED")
    parent.remove(el)


def insert_after(anchor, text):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)

    p = Paragraph(new_p, anchor._parent)

    try:
        p.style = anchor.style
    except Exception:
        pass

    p.add_run(text)
    return p


def one_contains(paragraphs, fragment, label):
    q = norm(fragment)

    hits = [
        (i, p)
        for i, p in enumerate(paragraphs, start=1)
        if q in norm(p.text)
    ]

    if len(hits) != 1:
        raise RuntimeError(
            f"{label}: expected exactly one match for "
            f"{fragment!r}; found {[i for i, _ in hits]}"
        )

    return hits[0]


def one_exact(paragraphs, text, label):
    q = norm(text)

    hits = [
        (i, p)
        for i, p in enumerate(paragraphs, start=1)
        if norm(p.text) == q
    ]

    if len(hits) != 1:
        raise RuntimeError(
            f"{label}: expected exactly one exact heading; "
            f"found {[i for i, _ in hits]}"
        )

    return hits[0]


def index_of(paragraphs, target):
    for i, p in enumerate(paragraphs):
        if p._p is target._p:
            return i
    raise RuntimeError("PARAGRAPH_OBJECT_NOT_FOUND")


# ======================================================================
# Verify input
# ======================================================================

print()
print("=" * 96)
print("STEP23C-QA2B1C: APPLY QA2E 01-14 LOCKED DECISIONS")
print("=" * 96)
print()

if not INPUT.exists():
    raise SystemExit(f"STOP: missing input:\n{INPUT}")

input_sha_before = sha256(INPUT)

print("INPUT:")
print(" ", INPUT.relative_to(ROOT))
print()
print("INPUT SHA256:")
print(" ", input_sha_before)
print()

doc = Document(INPUT)
initial = list(doc.paragraphs)

print("DOCX paragraphs before edit:", len(initial))
print()


# ======================================================================
# Resolve ALL targets BEFORE any modification
# ======================================================================

i004, p004 = one_contains(
    initial,
    "limited by cohort-specific noise. We asked whether a reproducible",
    "QA2C-004"
)

i006, p006 = one_contains(
    initial,
    "A conserved fibroblast CORE2 response was identified across cohorts",
    "QA2C-006"
)

i009, p009 = one_contains(
    initial,
    "These results nominate a conserved fibroblast CORE2 response as a",
    "QA2C-009"
)

i010, p010 = one_contains(
    initial,
    "reproducible treatment-associated stromal feature in ESCC",
    "QA2C-010"
)

i011, p011 = one_contains(
    initial,
    "In an independent neoadjuvant single-cell cohort (OMIX005710)",
    "QA2C-011"
)

i029, p029 = one_contains(
    initial,
    "Definitive ESCC-only external validation in OMIX005710",
    "QA2C-029"
)

i040, p040 = one_contains(
    initial,
    "The OMIX005710 analysis provides an important negative external finding",
    "QA2C-040"
)

i042, p042 = one_contains(
    initial,
    "A key limitation of the OMIX005710 analysis",
    "QA2C-042"
)

i048, p048 = one_contains(
    initial,
    "14. External validation in OMIX005710",
    "QA2C-048"
)

i049, p049 = one_contains(
    initial,
    "External validation was additionally assessed using the publicly available OMIX005710",
    "QA2C-049"
)

i050, p050 = one_contains(
    initial,
    "The source cohort contained one patient with esophageal adenosquamous carcinoma",
    "QA2C-050"
)

i051, p051 = one_contains(
    initial,
    "Supplementary Figure S2: Discordant / non-conserved signaling candidates",
    "QA2C-051"
)

i_da, p_da_heading = one_exact(
    initial,
    "Data Availability",
    "QA2C-052"
)

i_code, p_code_heading = one_exact(
    initial,
    "Code Availability",
    "QA2C-052 section boundary"
)

if i_code <= i_da:
    raise RuntimeError(
        "QA2C-052: Code Availability does not follow Data Availability."
    )

da_body = [
    p for p in initial[i_da:i_code - 1]
    if p.text.strip()
]

if not da_body:
    raise RuntimeError(
        "QA2C-052: no Data Availability body paragraph located."
    )

p_da_body = da_body[0]
extra_da_body = da_body[1:]


i_sm, p_sm_heading = one_exact(
    initial,
    "Supplementary Material Index",
    "QA2C-053"
)

# Prefer the existing aggregate supplementary-table index entry.
sm_anchor_hits = [
    (i, p)
    for i, p in enumerate(initial, start=1)
    if i > i_sm
    and "supplementary tables s1-s12" in norm(p.text)
]

if len(sm_anchor_hits) == 1:
    _, p_sm_anchor = sm_anchor_hits[0]
else:
    # Conservative fallback: use the last existing Supplementary*
    # entry after the heading.
    candidates = [
        p
        for i, p in enumerate(initial, start=1)
        if i > i_sm
        and norm(p.text).startswith("supplementary ")
    ]

    if not candidates:
        raise RuntimeError(
            "QA2C-053: unable to locate supplementary index anchor."
        )

    p_sm_anchor = candidates[-1]


# References must remain untouched.
i_ref, p_ref_heading = one_exact(
    initial,
    "References",
    "QA2C-054"
)

i_fig, p_fig_heading = one_exact(
    initial,
    "Figure Legends",
    "QA2C-054 references boundary"
)

if i_fig <= i_ref:
    raise RuntimeError(
        "QA2C-054: Figure Legends does not follow References."
    )

reference_text_before = [
    p.text
    for p in initial[i_ref:i_fig - 1]
]


# QA2C-009 + QA2C-010 are one approved logical sentence.
idx009 = index_of(initial, p009)
idx011 = index_of(initial, p011)

if idx011 <= idx009 + 1:
    raise RuntimeError(
        "QA2C-009/010: expected a multi-paragraph logical block."
    )

logical_block = initial[idx009:idx011]

if not any(p._p is p010._p for p in logical_block):
    raise RuntimeError(
        "QA2C-010 is not inside QA2C-009 combined logical block."
    )


# ======================================================================
# Prepare output
# ======================================================================

if OUTPUT.exists():
    OUTPUT.unlink()

shutil.copy2(INPUT, OUTPUT)

# Reload output so input remains physically untouched.
doc = Document(OUTPUT)
paras = list(doc.paragraphs)


# Re-resolve corresponding paragraph objects in OUTPUT
def out_contains(fragment, label):
    return one_contains(paras, fragment, label)[1]


p004 = out_contains(
    "limited by cohort-specific noise. We asked whether a reproducible",
    "QA2C-004"
)
p006 = out_contains(
    "A conserved fibroblast CORE2 response was identified across cohorts",
    "QA2C-006"
)
p009 = out_contains(
    "These results nominate a conserved fibroblast CORE2 response as a",
    "QA2C-009"
)
p010 = out_contains(
    "reproducible treatment-associated stromal feature in ESCC",
    "QA2C-010"
)
p011 = out_contains(
    "In an independent neoadjuvant single-cell cohort (OMIX005710)",
    "QA2C-011"
)
p029 = out_contains(
    "Definitive ESCC-only external validation in OMIX005710",
    "QA2C-029"
)
p040 = out_contains(
    "The OMIX005710 analysis provides an important negative external finding",
    "QA2C-040"
)
p042 = out_contains(
    "A key limitation of the OMIX005710 analysis",
    "QA2C-042"
)
p048 = out_contains(
    "14. External validation in OMIX005710",
    "QA2C-048"
)
p049 = out_contains(
    "External validation was additionally assessed using the publicly available OMIX005710",
    "QA2C-049"
)
p050 = out_contains(
    "The source cohort contained one patient with esophageal adenosquamous carcinoma",
    "QA2C-050"
)
p051 = out_contains(
    "Supplementary Figure S2: Discordant / non-conserved signaling candidates",
    "QA2C-051"
)

current = list(doc.paragraphs)

_, p_da_heading = one_exact(
    current, "Data Availability", "QA2C-052"
)
_, p_code_heading = one_exact(
    current, "Code Availability", "QA2C-052 boundary"
)

current = list(doc.paragraphs)
di = index_of(current, p_da_heading)
ci = index_of(current, p_code_heading)

da_body_out = [
    p for p in current[di + 1:ci]
    if p.text.strip()
]

if not da_body_out:
    raise RuntimeError("QA2C-052: no output Data Availability body.")

p_da_body = da_body_out[0]
extra_da_body = da_body_out[1:]

_, p_sm_heading = one_exact(
    list(doc.paragraphs),
    "Supplementary Material Index",
    "QA2C-053"
)

current = list(doc.paragraphs)
si = index_of(current, p_sm_heading)

anchors = [
    p
    for p in current[si + 1:]
    if "supplementary tables s1-s12" in norm(p.text)
]

if len(anchors) == 1:
    p_sm_anchor = anchors[0]
else:
    candidates = [
        p
        for p in current[si + 1:]
        if norm(p.text).startswith("supplementary ")
    ]
    if not candidates:
        raise RuntimeError(
            "QA2C-053: output supplementary anchor missing."
        )
    p_sm_anchor = candidates[-1]


audit = []


def record(item_id, action, before, after):
    audit.append({
        "item_id": item_id,
        "action": action,
        "before": before,
        "after": after,
    })


# ======================================================================
# Apply QA2E 01-14
# ======================================================================

# QA2C-004 — replace only the requested sentence tail.
before = p004.text
marker = "We asked whether"
pos = before.find(marker)

if pos < 0:
    raise RuntimeError(
        "QA2C-004: sentence marker not found."
    )

after = before[:pos] + A004
set_paragraph_text(p004, after)
record("QA2C-004", "CONTEXTUAL_SENTENCE_REWRITE", before, after)


# QA2C-006
before = p006.text
set_paragraph_text(p006, A006)
record("QA2C-006", "REPLACE_SENTENCE_PARAGRAPH", before, A006)


# QA2C-009 + QA2C-010 — one logical sentence spanning multiple DOCX paras.
current = list(doc.paragraphs)
s = index_of(current, p009)
e = index_of(current, p011)

group = current[s:e]

if len(group) < 2:
    raise RuntimeError(
        "QA2C-009/010: combined logical block unexpectedly short."
    )

before_combined = " ".join(
    p.text.strip()
    for p in group
    if p.text.strip()
)

set_paragraph_text(group[0], A009_010)

for p in group[1:]:
    delete_paragraph(p)

record(
    "QA2C-009|QA2C-010",
    "REPLACE_COMBINED_LOGICAL_SENTENCE",
    before_combined,
    A009_010
)


# QA2C-011
before = p011.text
set_paragraph_text(p011, A011)
record("QA2C-011", "REPLACE_FULL_PARAGRAPH", before, A011)


# QA2C-029
before = p029.text
set_paragraph_text(p029, A029)
record("QA2C-029", "REPLACE_FULL_PARAGRAPH", before, A029)


# QA2C-040
before = p040.text
set_paragraph_text(p040, A040)
record("QA2C-040", "REPLACE_FULL_PARAGRAPH", before, A040)


# QA2C-042
before = p042.text
set_paragraph_text(p042, A042)
record("QA2C-042", "REPLACE_FULL_PARAGRAPH", before, A042)


# QA2C-048
before = p048.text
set_paragraph_text(p048, A048)
record("QA2C-048", "REPLACE_HEADING", before, A048)


# QA2C-049
before = p049.text
set_paragraph_text(p049, A049)
record("QA2C-049", "REPLACE_FULL_PARAGRAPH", before, A049)


# QA2C-050
before = p050.text
set_paragraph_text(p050, A050)
record("QA2C-050", "REPLACE_FULL_PARAGRAPH", before, A050)


# QA2C-051
before = p051.text
set_paragraph_text(p051, A051)
record("QA2C-051", "REPLACE_INDEX_ENTRY", before, A051)


# QA2C-052 — Data Availability body.
before = " ".join(
    p.text.strip()
    for p in [p_da_body] + extra_da_body
    if p.text.strip()
)

set_paragraph_text(p_da_body, A052)

for p in extra_da_body:
    delete_paragraph(p)

record(
    "QA2C-052",
    "REPLACE_DATA_AVAILABILITY_BODY",
    before,
    A052
)


# QA2C-053 — append frozen supplementary audit entries.
existing_all = "\n".join(
    p.text for p in doc.paragraphs
)

for x in A053:
    if norm(x) in norm(existing_all):
        raise RuntimeError(
            "QA2C-053: approved supplementary entry already exists: "
            + x
        )

anchor = p_sm_anchor

for x in A053:
    anchor = insert_after(anchor, x)

record(
    "QA2C-053",
    "APPEND_SUPPLEMENTARY_INDEX_ENTRIES",
    "",
    "\n".join(A053)
)


# QA2C-054 — deliberate NO-EDIT hold.
record(
    "QA2C-054",
    "HOLD_FOR_BIBLIOGRAPHIC_VERIFICATION",
    "<References unchanged>",
    "<NO TEXT INSERTED>"
)


# ======================================================================
# Structural/scientific QA BEFORE SAVE
# ======================================================================

final_text = "\n".join(
    p.text for p in doc.paragraphs
)

required = [
    A004,
    A006,
    A009_010,
    A011,
    A029,
    A040,
    A042,
    A048,
    A049,
    A050,
    A051,
    A052,
] + A053

missing = [
    x for x in required
    if norm(x) not in norm(final_text)
]

if missing:
    print()
    print("STOP BEFORE SAVE: approved text missing:")
    for x in missing:
        print(" ", x)
    OUTPUT.unlink(missing_ok=True)
    raise SystemExit(
        "QA2E BATCH NOT SAVED."
    )


legacy_forbidden = [
    "A conserved fibroblast CORE2 response was identified across cohorts",
    "The OMIX005710 analysis provides an important negative external finding",
    "14. External validation in OMIX005710",
    "Supplementary Figure S2: Discordant / non-conserved signaling candidates",
]

legacy_hits = [
    x for x in legacy_forbidden
    if norm(x) in norm(final_text)
]

if legacy_hits:
    print()
    print("STOP BEFORE SAVE: legacy locked wording remains:")
    for x in legacy_hits:
        print(" ", x)
    OUTPUT.unlink(missing_ok=True)
    raise SystemExit(
        "QA2E LEGACY-WORDING CHECK FAILED."
    )


# References must be byte-for-text unchanged within section.
current = list(doc.paragraphs)

_, ref_h = one_exact(
    current,
    "References",
    "QA2C-054 post-edit"
)

_, fig_h = one_exact(
    current,
    "Figure Legends",
    "QA2C-054 post-edit boundary"
)

ri = index_of(current, ref_h)
fi = index_of(current, fig_h)

reference_text_after = [
    p.text
    for p in current[ri + 1:fi]
]

if reference_text_after != reference_text_before:
    OUTPUT.unlink(missing_ok=True)
    raise SystemExit(
        "CRITICAL STOP: References section changed despite QA2C-054 HOLD."
    )


# ======================================================================
# Save
# ======================================================================

doc.save(OUTPUT)


# Input immutability
if sha256(INPUT) != input_sha_before:
    raise SystemExit(
        "CRITICAL STOP: QA2B1B source DOCX was modified."
    )


# Reopen for validity
check = Document(OUTPUT)

if len(check.paragraphs) < 1:
    raise SystemExit(
        "CRITICAL STOP: output DOCX invalid."
    )


# ======================================================================
# Audit + hold file + checksums
# ======================================================================

OUT_AUDIT.parent.mkdir(
    parents=True,
    exist_ok=True
)

with open(
    OUT_AUDIT,
    "w",
    encoding="utf-8-sig",
    newline=""
) as f:

    w = csv.DictWriter(
        f,
        fieldnames=[
            "item_id",
            "action",
            "before",
            "after",
        ]
    )

    w.writeheader()
    w.writerows(audit)


OUT_HOLD.write_text(
    "\n".join([
        "STEP23C-QA2B1C REFERENCE VERIFICATION HOLD",
        "=" * 72,
        "",
        "QA2C-054",
        "Decision: REQUIRE_VERIFIED_CITATIONS",
        "Execution mode: HOLD_FOR_BIBLIOGRAPHIC_VERIFICATION",
        "",
        "NO bibliography text was generated or inserted by QA2B1C.",
        "",
        "Required later verification:",
        "  - GSE197677 source publication",
        "  - GSE221561 source publication",
        "  - OMIX005710 source publication / primary repository record",
        "",
        "Author names, article title, journal, year, DOI and final citation",
        "format must be checked against primary publication records before",
        "insertion.",
        "",
    ]),
    encoding="utf-8"
)


with open(
    OUT_SHA,
    "w",
    encoding="utf-8"
) as f:

    for p in [
        OUTPUT,
        OUT_AUDIT,
        OUT_HOLD,
    ]:
        f.write(
            f"{sha256(p)}  {p.relative_to(ROOT)}\n"
        )


# ======================================================================
# Summary
# ======================================================================

print()
print("=" * 96)
print("STEP23C-QA2B1C COMPLETE")
print("=" * 96)
print()

print("QA2B1B SOURCE MODIFIED          : NO")
print("STEP22A RECOMPUTED              : NO")
print("GLOBAL FIND/REPLACE             : NO")
print()

print("QA2E ITEMS                      : 15")
print("QA2E TEXT/PACKAGE ITEMS APPLIED : 14")
print("REFERENCE VERIFICATION HOLD     : 1")
print()

print("APPLIED / LOCKED ACTIONS:")
for r in audit:
    print(
        f"  {r['item_id']:<19s} {r['action']}"
    )

print()
print("OUTPUT DOCX:")
print(" ", OUTPUT.relative_to(ROOT))

print()
print("AUDIT:")
print(" ", OUT_AUDIT.relative_to(ROOT))

print()
print("REFERENCE HOLD:")
print(" ", OUT_HOLD.relative_to(ROOT))

print()
print("STATUS:")
print(
    "STEP23C_QA2B1C_QA2E_01_14_APPLY_PASS_REFERENCE_HOLD"
)

print()
print("NEXT:")
print(
    "RENDER THE NEW DOCX AND PERFORM FULL VISUAL + TEXT QA "
    "BEFORE ANY REFERENCE INSERTION."
)
print()

