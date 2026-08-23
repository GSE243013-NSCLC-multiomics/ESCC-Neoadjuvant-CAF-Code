#!/usr/bin/env python3

from pathlib import Path
import re
import hashlib
import shutil

from docx import Document

PROJECT = Path.cwd()

INFILE = (
    PROJECT
    / "04_results"
    / "manuscript_update"
    / "Step23"
    / "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_WORKING.docx"
)

OUTFILE = (
    PROJECT
    / "04_results"
    / "manuscript_update"
    / "Step23"
    / "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA1B.docx"
)

STATUS = (
    PROJECT
    / "00_metadata"
    / "Step23_manuscript_update"
    / "STEP23C_QA1B_PLACEMENT_STATUS.txt"
)

ABSTRACT_INSERT = (
    "In an independent neoadjuvant single-cell cohort (OMIX005710), "
    "definitive ESCC-only validation was precluded because the single "
    "patient with esophageal adenosquamous carcinoma could not be "
    "identified from the available patient-level metadata. In a "
    "pre-specified secondary histology-sensitivity analysis spanning all "
    "22 possible EASC assignments, the treatment-associated CORE "
    "direction was consistently opposite to that observed in the "
    "discovery cohorts."
)

METHODS_1 = (
    "External validation was additionally assessed using the publicly "
    "available OMIX005710 neoadjuvant single-cell RNA-sequencing cohort. "
    "Fibroblast pseudobulk profiles were constructed from raw counts after "
    "broad cell-type annotation using a marker framework independent of "
    "the CORE genes. A pre-specified minimum of 20 fibroblast cells per "
    "patient-timepoint was required for technical eligibility; one sample "
    "(P15 before treatment, 18 fibroblasts) therefore remained in raw "
    "provenance but was excluded from the technical reference and paired "
    "analysis. Raw pseudobulk counts were transformed to library-size CPM "
    "and log2(CPM+1), followed by gene-wise z-scoring within the eligible "
    "OMIX tumor reference. CORE2_STRONG and CORE4 were calculated using "
    "the previously frozen formulas without refitting gene weights."
)

METHODS_2 = (
    "The source cohort contained one patient with esophageal "
    "adenosquamous carcinoma (EASC), but the corresponding individual "
    "patient could not be resolved from the publicly available metadata. "
    "Accordingly, definitive ESCC-only primary validation was considered "
    "unavailable. Before examining CORE scores, we therefore pre-specified "
    "a secondary histology-sensitivity framework comprising all 22 "
    "possible assignments of the EASC identity. Each scenario excluded "
    "the corresponding patient's technically eligible tumor observations "
    "from the z-score reference and, when applicable, from the paired "
    "analysis. Paired treatment effects were defined as after-treatment "
    "minus before-treatment score. Robustness to individual paired "
    "patients was assessed by full leave-one-patient-out recomputation, "
    "with the omitted patient removed from both the z-score reference and "
    "the paired effect calculation."
)

METHOD_HEADING = "14. External validation in OMIX005710"


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def move_before(paragraph, target):
    # Moving the existing XML node preserves paragraph formatting.
    target._p.addprevious(paragraph._p)


def find_exact(doc, text):
    target = norm(text)
    hits = [
        p for p in doc.paragraphs
        if norm(p.text) == target
    ]
    if len(hits) != 1:
        raise RuntimeError(
            f"Expected exactly one paragraph match, found {len(hits)}:\n"
            f"{text[:120]}"
        )
    return hits[0]


def find_regex(doc, pattern):
    rx = re.compile(pattern, re.I)
    hits = [
        p for p in doc.paragraphs
        if rx.search(norm(p.text))
    ]
    if len(hits) != 1:
        raise RuntimeError(
            f"Expected exactly one regex match, found {len(hits)}: {pattern}"
        )
    return hits[0]


def find_optional(doc, pattern):
    rx = re.compile(pattern, re.I)
    hits = [
        p for p in doc.paragraphs
        if rx.search(norm(p.text))
    ]
    if len(hits) == 1:
        return hits[0]
    if len(hits) == 0:
        return None
    raise RuntimeError(
        f"Expected 0 or 1 regex match, found {len(hits)}: {pattern}"
    )


def paragraph_index(doc, target):
    for i, p in enumerate(doc.paragraphs):
        if p._p is target._p:
            return i
    raise RuntimeError("Paragraph not found in document")


if not INFILE.exists():
    raise SystemExit(f"Missing input DOCX:\n{INFILE}")

source_sha_before = sha256(INFILE)

if OUTFILE.exists():
    OUTFILE.unlink()

shutil.copy2(INFILE, OUTFILE)

doc = Document(OUTFILE)

print("=" * 72)
print("STEP23C-QA1B: FIX INSERTION PLACEMENT")
print("=" * 72)
print()

# ============================================================
# 1. FIX ABSTRACT
# ============================================================

abstract_insert = find_exact(doc, ABSTRACT_INSERT)

wordcount_p = find_regex(
    doc,
    r"\(\s*Approximately\s+\d+\s+words?\s*\.\s*\)"
)

# Move OMIX sentence before the abstract word-count line.
move_before(
    abstract_insert,
    wordcount_p
)

print("✓ Abstract OMIX paragraph moved before word-count line")

# Recompute Abstract word count.
abstract_heading = find_regex(
    doc,
    r"^\s*Abstract\s*$"
)

intro_heading = find_regex(
    doc,
    r"^\s*(?:\d+\s*[\.\)]?\s*)?Introduction\s*$"
)

a_idx = paragraph_index(
    doc,
    abstract_heading
)

i_idx = paragraph_index(
    doc,
    intro_heading
)

abstract_parts = []

for p in doc.paragraphs[a_idx + 1:i_idx]:
    t = norm(p.text)

    if not t:
        continue

    if p._p is wordcount_p._p:
        continue

    abstract_parts.append(t)

abstract_text = " ".join(abstract_parts)

abstract_words = re.findall(
    r"\b[\w][\w'\-/+]*\b",
    abstract_text
)

wordcount_p.text = (
    f"(Approximately {len(abstract_words)} words.)"
)

print(
    "✓ Abstract word-count marker updated:",
    len(abstract_words),
    "words"
)

# ============================================================
# 2. FIX METHODS PLACEMENT
# ============================================================

methods1 = find_exact(
    doc,
    METHODS_1
)

methods2 = find_exact(
    doc,
    METHODS_2
)

data_availability = find_regex(
    doc,
    r"^\s*Data\s+Availability\s*$"
)

# Remove a previously created QA1B heading if script is ever rerun
# on a derivative file.
old_heading_hits = [
    p for p in doc.paragraphs
    if norm(p.text) == METHOD_HEADING
]

for p in old_heading_hits:
    remove_paragraph(p)

# Re-resolve paragraphs after possible XML changes.
methods1 = find_exact(
    doc,
    METHODS_1
)

methods2 = find_exact(
    doc,
    METHODS_2
)

data_availability = find_regex(
    doc,
    r"^\s*Data\s+Availability\s*$"
)

# Move the two methods paragraphs into Methods,
# immediately before Data Availability.
move_before(
    methods1,
    data_availability
)

move_before(
    methods2,
    data_availability
)

# Find numbered Methods item 13 to copy its heading style.
method13_candidates = [
    p for p in doc.paragraphs
    if re.match(
        r"^\s*13\s*[\.\)]",
        norm(p.text)
    )
]

if len(method13_candidates) != 1:
    raise RuntimeError(
        f"Expected exactly one Methods item 13; "
        f"found {len(method13_candidates)}"
    )

method13 = method13_candidates[0]

# Re-find methods1 after move.
methods1 = find_exact(
    doc,
    METHODS_1
)

new_heading = methods1.insert_paragraph_before(
    METHOD_HEADING
)

try:
    new_heading.style = method13.style
except Exception:
    pass

print(
    "✓ OMIX Methods moved before Data Availability"
)

print(
    "✓ Added Methods subheading:",
    METHOD_HEADING
)

# ============================================================
# 3. ORDER CHECKS
# ============================================================

abstract_insert = find_exact(
    doc,
    ABSTRACT_INSERT
)

wordcount_p = find_regex(
    doc,
    r"\(\s*Approximately\s+\d+\s+words?\s*\.\s*\)"
)

intro_heading = find_regex(
    doc,
    r"^\s*(?:\d+\s*[\.\)]?\s*)?Introduction\s*$"
)

methods1 = find_exact(
    doc,
    METHODS_1
)

methods2 = find_exact(
    doc,
    METHODS_2
)

method14 = find_exact(
    doc,
    METHOD_HEADING
)

data_availability = find_regex(
    doc,
    r"^\s*Data\s+Availability\s*$"
)

acknowledgements = find_optional(
    doc,
    r"^\s*(?:\d+\s*[\.\)]?\s*)?Acknowledg(e)?ments?\s*$"
)

checks = {
    "abstract_insert_before_wordcount":
        paragraph_index(doc, abstract_insert)
        <
        paragraph_index(doc, wordcount_p),

    "wordcount_before_introduction":
        paragraph_index(doc, wordcount_p)
        <
        paragraph_index(doc, intro_heading),

    "method14_before_methods_text":
        paragraph_index(doc, method14)
        <
        paragraph_index(doc, methods1),

    "methods1_before_methods2":
        paragraph_index(doc, methods1)
        <
        paragraph_index(doc, methods2),

    "methods2_before_data_availability":
        paragraph_index(doc, methods2)
        <
        paragraph_index(doc, data_availability),
}

if acknowledgements is not None:
    checks["methods_block_before_acknowledgements"] = (
        paragraph_index(doc, methods2)
        <
        paragraph_index(doc, acknowledgements)
    )

for name, passed in checks.items():
    print(
        f"{name:<42s}:",
        "PASS" if passed else "FAIL"
    )

if not all(checks.values()):
    raise RuntimeError(
        "Placement QA failed"
    )

# ============================================================
# 4. SAVE
# ============================================================

doc.save(OUTFILE)

source_sha_after = sha256(INFILE)
output_sha = sha256(OUTFILE)

if source_sha_before != source_sha_after:
    raise RuntimeError(
        "CRITICAL: STEP23C_WORKING source changed"
    )

STATUS.write_text(
    "\n".join([
        "STEP23C-QA1B PLACEMENT FIX",
        "",
        f"Input: {INFILE}",
        f"Input SHA256: {source_sha_before}",
        "",
        f"Output: {OUTFILE}",
        f"Output SHA256: {output_sha}",
        "",
        f"Abstract words: {len(abstract_words)}",
        "",
        "Abstract OMIX placement: PASS",
        "Methods OMIX placement: PASS",
        "Original STEP23C working copy modified: NO",
        "",
        "STATUS:",
        "PLACEMENT_FIX_COMPLETE_READY_FOR_RERENDER",
        ""
    ]),
    encoding="utf-8"
)

print()
print("=" * 72)
print("STEP23C-QA1B COMPLETE")
print("=" * 72)
print()

print(
    "ORIGINAL STEP23C WORKING MODIFIED : NO"
)

print(
    "ABSTRACT OMIX POSITION            : FIXED"
)

print(
    "ABSTRACT WORD COUNT               :",
    len(abstract_words)
)

print(
    "METHODS OMIX POSITION             : FIXED"
)

print(
    "NEW METHODS SUBHEADING            :",
    METHOD_HEADING
)

print()
print("OUTPUT:")
print(OUTFILE)
print()

print(
    "STATUS:"
)

print(
    "PLACEMENT_FIX_COMPLETE_READY_FOR_RERENDER"
)
