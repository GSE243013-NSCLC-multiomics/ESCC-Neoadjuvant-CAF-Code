#!/usr/bin/env python3

from pathlib import Path
import csv
import hashlib
import shutil
from docx import Document

ROOT = Path.cwd()

META = ROOT / "00_metadata" / "Step23_manuscript_update"
RESULT = ROOT / "04_results" / "manuscript_update" / "Step23"

QA1B = RESULT / "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA1B.docx"

INPUT = (
    RESULT /
    "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1A_SAFE_EDITS.docx"
)

OUTPUT = (
    RESULT /
    "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1B_CONTEXTUAL_EDITS.docx"
)

MAP_CSV = META / "STEP23C_QA2B0_CONTEXT_MAP.csv"

OUT_AUDIT = META / "STEP23C_QA2B1B_EDIT_AUDIT.csv"
OUT_SHA = META / "STEP23C_QA2B1B_CHECKSUMS.sha256"

EXPECTED_QA1B_SHA = (
    "0e91d04f82c3a3353bc6214b94312f72767f7da50778d40d0f33c016bc65f598"
)


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def run_signature(run):
    size = run.font.size.pt if run.font.size else None

    return (
        run.style.name if run.style else None,
        run.bold,
        run.italic,
        run.underline,
        run.font.name,
        size,
    )


def replace_once_preserve_runs(paragraph, old, new):
    runs = paragraph.runs

    if not runs:
        return False, "NO_RUNS"

    full = "".join(r.text for r in runs)

    start = full.find(old)

    if start < 0:
        return False, "TEXT_NOT_FOUND"

    if full.find(old, start + 1) >= 0:
        return False, "MULTIPLE_MATCHES"

    stop = start + len(old)

    positions = []
    cursor = 0

    for i, run in enumerate(runs):
        a = cursor
        b = cursor + len(run.text)
        positions.append((i, a, b))
        cursor = b

    touched = [
        i for i, a, b in positions
        if b > start and a < stop
    ]

    if not touched:
        return False, "RUN_MAPPING_FAILED"

    sigs = {
        run_signature(runs[i])
        for i in touched
        if runs[i].text
    }

    if len(sigs) > 1:
        return False, "MIXED_RUN_FORMATTING"

    first_i = touched[0]
    last_i = touched[-1]

    first_global = positions[first_i][1]
    last_global = positions[last_i][1]

    local_start = start - first_global
    local_stop = stop - last_global

    expected = full[:start] + new + full[stop:]

    if first_i == last_i:
        r = runs[first_i]
        r.text = (
            r.text[:local_start]
            + new
            + r.text[local_start + len(old):]
        )

    else:
        first = runs[first_i]
        last = runs[last_i]

        prefix = first.text[:local_start]
        suffix = last.text[local_stop:]

        first.text = prefix + new

        for i in touched[1:-1]:
            runs[i].text = ""

        last.text = suffix

    actual = "".join(r.text for r in runs)

    if actual != expected:
        raise RuntimeError(
            "POST_REPLACEMENT_TEXT_MISMATCH"
        )

    return True, "APPLIED"


def sentence_containing(text, anchor):
    pos = text.find(anchor)

    if pos < 0:
        return None

    prev = text.rfind(". ", 0, pos)

    start = 0 if prev < 0 else prev + 2

    end = text.find(".", pos)

    if end < 0:
        end = len(text)
    else:
        end += 1

    return text[start:end]


print()
print("=" * 92)
print("STEP23C-QA2B1B V2: APPLY CONTEXTUAL LOCKED EDITS")
print("=" * 92)
print()

for p in [QA1B, INPUT, MAP_CSV]:
    if not p.exists():
        raise SystemExit(
            f"STOP: missing required input: {p}"
        )

if sha256(QA1B) != EXPECTED_QA1B_SHA:
    raise SystemExit(
        "STOP: frozen QA1B manuscript SHA256 changed."
    )

print("✓ Frozen QA1B verified.")
print()

maps = {
    r["item_id"].strip(): r
    for r in read_csv(MAP_CSV)
}

# Always rebuild from the successful QA2B1A output.
if OUTPUT.exists():
    OUTPUT.unlink()

shutil.copy2(INPUT, OUTPUT)

doc = Document(OUTPUT)

before = [p.text for p in doc.paragraphs]

audit = []


# ============================================================
# EDITS
#
# Important corrections in V2:
#
# QA2C-017 / 018 are consecutive DOCX paragraphs 67 / 68.
# They therefore remain two paragraph edits.
#
# QA2C-032 / 033 are likewise consecutive paragraphs 124 / 125.
#
# QA2C-030 / 031 are duplicate risk hits on one actual paragraph,
# therefore one actual manuscript edit.
# ============================================================

EDITS = [

    # --------------------------------------------------------
    # Abstract contextual validation wording
    # --------------------------------------------------------
    {
        "ids": ["QA2C-005"],
        "mode": "SENTENCE",
        "anchor":
            "Orthogonal regulator activity proxies",
        "new":
            "Orthogonal regulator activity proxies and a curated "
            "ligand–receptor resource were used for focused "
            "follow-up evaluation.",
    },

    # --------------------------------------------------------
    # Introduction — sentence split across two DOCX paragraphs
    # --------------------------------------------------------
    {
        "ids": ["QA2C-017"],
        "mode": "FRAGMENT",
        "old":
            "Our objective was to define a reproducible fibroblast "
            "response across",
        "new":
            "Our objective was to define a fibroblast response "
            "reproducible across the two discovery",
    },

    {
        "ids": ["QA2C-018"],
        "mode": "FRAGMENT",
        "old":
            "cohorts and to separate strong, conserved effects "
            "(the CORE2 module) from",
        "new":
            "cohorts and to distinguish the stronger "
            "discovery-cohort-reproducible effects represented "
            "by the CORE2 module from",
    },

    # --------------------------------------------------------
    # Results
    # --------------------------------------------------------
    {
        "ids": ["QA2C-021"],
        "mode": "SENTENCE",
        "anchor":
            "conservation was the strongest and most reproducible "
            "project-level result",
        "new":
            "Across the two discovery cohorts, the CORE2 response "
            "was the strongest and most reproducible project-level "
            "treatment-associated result.",
    },

    {
        "ids": ["QA2C-027"],
        "mode": "FRAGMENT",
        "old":
            "association directions across cohorts and are not conserved.",
        "new":
            "association directions across the two discovery cohorts "
            "and were not conserved between them.",
    },

    # --------------------------------------------------------
    # Discussion — duplicate audit hit, one actual edit
    # --------------------------------------------------------
    {
        "ids": ["QA2C-030", "QA2C-031"],
        "mode": "SENTENCE",
        "anchor":
            "A conserved fibroblast CDKN1B",
        "new":
            "Across the two discovery cohorts, the fibroblast "
            "CDKN1B–GSN CORE2 response was the most reproducible "
            "treatment-associated fibroblast result.",
    },

    # --------------------------------------------------------
    # Discussion — sentence split over paragraphs 124 / 125
    # --------------------------------------------------------
    {
        "ids": ["QA2C-032"],
        "mode": "FRAGMENT",
        "old":
            "response is more reproducible than any individual "
            "upstream ligand or",
        "new":
            "Within the two discovery cohorts, the CORE2 response "
            "was more reproducible than any individual upstream "
            "ligand or",
    },

    {
        "ids": ["QA2C-033"],
        "mode": "FRAGMENT",
        "old":
            "regulator mechanism identified in this project.",
        "new":
            "regulator candidate evaluated in this project.",
    },

    # --------------------------------------------------------
    # Discussion
    # --------------------------------------------------------
    {
        "ids": ["QA2C-039"],
        "mode": "SENTENCE",
        "anchor":
            "The conserved fibroblast CORE2 response is the stable "
            "result of this project",
        "new":
            "Within the two discovery cohorts, the fibroblast CORE2 "
            "response was the most stable treatment-associated result "
            "of this project; upstream ligand and regulator evidence "
            "is weaker and should be treated as hypothesis-generating.",
    },

]


for edit in EDITS:

    ids = edit["ids"]

    # For a true duplicate group, all IDs must map to the same
    # paragraph. Single-ID entries simply use their own mapping.
    indices = []

    for item_id in ids:

        if item_id not in maps:
            raise SystemExit(
                f"STOP: {item_id} missing from context map."
            )

        x = (
            maps[item_id]
            .get("docx_paragraph_index1", "")
            .strip()
        )

        if not x.isdigit():
            raise SystemExit(
                f"STOP: no paragraph mapping for {item_id}"
            )

        indices.append(int(x))

    unique_indices = sorted(set(indices))

    if len(unique_indices) != 1:
        raise SystemExit(
            f"STOP: edit group {ids} mapped to "
            f"different paragraphs: {unique_indices}"
        )

    idx1 = unique_indices[0]

    if not (1 <= idx1 <= len(doc.paragraphs)):
        raise SystemExit(
            f"STOP: paragraph index out of range: {idx1}"
        )

    p = doc.paragraphs[idx1 - 1]

    before_text = p.text

    old = None

    if edit["mode"] == "SENTENCE":

        old = sentence_containing(
            before_text,
            edit["anchor"]
        )

        if old is None:
            ok = False
            reason = "ANCHOR_NOT_FOUND"

        else:
            ok, reason = replace_once_preserve_runs(
                p,
                old,
                edit["new"]
            )

    elif edit["mode"] == "FRAGMENT":

        old = edit["old"]

        ok, reason = replace_once_preserve_runs(
            p,
            old,
            edit["new"]
        )

    else:
        raise RuntimeError(
            "UNKNOWN_EDIT_MODE"
        )

    audit.append({
        "item_ids": "|".join(ids),
        "paragraph": idx1,
        "mode": edit["mode"],
        "status": "APPLIED" if ok else "SKIPPED",
        "reason": reason,
        "old_text": old or "",
        "new_text": edit["new"],
        "paragraph_before": before_text,
        "paragraph_after": p.text,
    })


# Never save a partially applied contextual batch.
skipped_rows = [
    r for r in audit
    if r["status"] != "APPLIED"
]

if skipped_rows:
    print()
    print("STOP BEFORE SAVE:")
    for r in skipped_rows:
        print(
            f"  {r['item_ids']}: "
            f"{r['reason']} "
            f"(paragraph {r['paragraph']})"
        )

    OUTPUT.unlink(missing_ok=True)

    raise SystemExit(
        "CONTEXTUAL BATCH NOT SAVED: "
        "one or more guarded edits did not match."
    )


doc.save(OUTPUT)


# ============================================================
# Structural QA
# ============================================================

check = Document(OUTPUT)

after = [p.text for p in check.paragraphs]

if len(after) != len(before):
    raise SystemExit(
        "CRITICAL STOP: DOCX paragraph count changed."
    )

changed = {
    i + 1
    for i, (a, b) in enumerate(
        zip(before, after)
    )
    if a != b
}

expected_changed = {
    int(r["paragraph"])
    for r in audit
}

unexpected = changed - expected_changed

if unexpected:
    raise SystemExit(
        "CRITICAL STOP: unexpected paragraphs changed: "
        + repr(sorted(unexpected))
    )

if changed != expected_changed:
    raise SystemExit(
        "CRITICAL STOP: expected edited paragraph set "
        "does not equal actual changed paragraph set."
    )


# ============================================================
# Scientific wording regression checks
# ============================================================

all_text = "\n".join(after)

must_be_absent = [
    "Our objective was to define a reproducible fibroblast response across",
    "A conserved fibroblast CDKN1B",
    "regulator mechanism identified in this project.",
    "The conserved fibroblast CORE2 response is the stable result of this project",
]

legacy_hits = [
    x for x in must_be_absent
    if x in all_text
]

if legacy_hits:
    print()
    print("STOP: legacy wording remains:")
    for x in legacy_hits:
        print(" ", x)

    raise SystemExit(
        "SCIENTIFIC WORDING REGRESSION CHECK FAILED."
    )


# ============================================================
# Audit
# ============================================================

fields = [
    "item_ids",
    "paragraph",
    "mode",
    "status",
    "reason",
    "old_text",
    "new_text",
    "paragraph_before",
    "paragraph_after",
]

with open(
    OUT_AUDIT,
    "w",
    encoding="utf-8-sig",
    newline=""
) as f:

    w = csv.DictWriter(
        f,
        fieldnames=fields
    )

    w.writeheader()
    w.writerows(audit)


# ============================================================
# Source immutability and checksums
# ============================================================

if sha256(QA1B) != EXPECTED_QA1B_SHA:
    raise SystemExit(
        "CRITICAL STOP: frozen QA1B was modified."
    )

with open(
    OUT_SHA,
    "w",
    encoding="utf-8"
) as f:

    for p in [OUTPUT, OUT_AUDIT]:
        f.write(
            f"{sha256(p)}  {p.relative_to(ROOT)}\n"
        )


# ============================================================
# Console summary
# ============================================================

print()
print("=" * 92)
print("STEP23C-QA2B1B V2 COMPLETE")
print("=" * 92)
print()

print("QA1B MODIFIED                 : NO")
print("STEP22A RECOMPUTED            : NO")
print("GLOBAL FIND/REPLACE           : NO")
print()

print("CONTEXTUAL EDIT ACTIONS       :", len(EDITS))
print("EDIT ACTIONS APPLIED          :", len(audit))
print("EDIT ACTIONS SKIPPED          : 0")
print("CHANGED DOCX PARAGRAPHS       :", len(changed))
print()

for r in audit:
    print(
        f"{r['item_ids']:<23s} "
        f"APPLIED  "
        f"paragraph={r['paragraph']}"
    )

print()
print("OUTPUT DOCX:")
print(" ", OUTPUT.relative_to(ROOT))
print()

print("AUDIT:")
print(" ", OUT_AUDIT.relative_to(ROOT))
print()

print("STATUS:")
print(
    "STEP23C_QA2B1B_CONTEXTUAL_EDIT_PASS"
)

print()
print("NEXT:")
print(
    "APPLY QA2E MUST-REVIEW DECISIONS AND "
    "COMPLETE PACKAGE-LEVEL ITEMS."
)
print()

