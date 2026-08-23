#!/usr/bin/env python3

from pathlib import Path
import csv
import hashlib
import shutil
from docx import Document

ROOT = Path.cwd()

META = ROOT / "00_metadata" / "Step23_manuscript_update"
RESULT = ROOT / "04_results" / "manuscript_update" / "Step23"

QA1B = RESULT / "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA1B.docx"
WORKING = RESULT / "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B_WORKING.docx"

QA2G_LOCK = META / "STEP23C_QA2G_REMAINING_DECISION_LOCK.csv"
CONTEXT_MAP = META / "STEP23C_QA2B0_CONTEXT_MAP.csv"

OUT_DOCX = RESULT / "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_QA2B1A_SAFE_EDITS.docx"
OUT_CSV = META / "STEP23C_QA2B1A_EDIT_AUDIT.csv"
OUT_QUEUE = RESULT / "STEP23C_QA2B1A_REMAINING_MANUAL_QUEUE.txt"
OUT_SHA = META / "STEP23C_QA2B1A_CHECKSUMS.sha256"

EXPECTED_QA1B_SHA = (
    "0e91d04f82c3a3353bc6214b94312f72767f7da50778d40d0f33c016bc65f598"
)

# These QA2G decisions were explicitly approved as
# non-contextual full-paragraph / fragment / heading replacements.
SAFE_IDS = {
    "QA2C-001",
    "QA2C-002",
    "QA2C-003",
    "QA2C-013",
    "QA2C-014",
    "QA2C-016",
    "QA2C-020",
    "QA2C-022",
    "QA2C-023",
    "QA2C-038",
    "QA2C-044",
    "QA2C-045",
    "QA2C-046",
    "QA2C-047",
}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def run_signature(run):
    size = run.font.size.pt if run.font.size else None

    return (
        run.style.name if run.style else None,
        run.bold,
        run.italic,
        run.underline,
        run.font.name,
        size,
    )


def replace_once_preserve_runs(paragraph, old, new):
    """
    Exact single replacement in paragraph.runs.

    If the match crosses runs with different formatting,
    refuse to edit rather than flatten formatting.
    """

    runs = paragraph.runs

    if not runs:
        return False, "NO_RUNS"

    full = "".join(r.text for r in runs)

    start = full.find(old)

    if start < 0:
        return False, "EXACT_TEXT_NOT_FOUND"

    if full.find(old, start + 1) >= 0:
        return False, "MULTIPLE_EXACT_MATCHES"

    stop = start + len(old)

    positions = []
    cursor = 0

    for i, run in enumerate(runs):
        a = cursor
        b = cursor + len(run.text)
        positions.append((i, a, b))
        cursor = b

    touched = []

    for i, a, b in positions:
        if b <= start:
            continue
        if a >= stop:
            continue
        touched.append(i)

    if not touched:
        return False, "RUN_MAPPING_FAILED"

    sigs = {
        run_signature(runs[i])
        for i in touched
        if runs[i].text
    }

    if len(sigs) > 1:
        return False, "MIXED_RUN_FORMATTING"

    first_i = touched[0]
    last_i = touched[-1]

    first_start = positions[first_i][1]
    last_start = positions[last_i][1]

    local_start = start - first_start
    local_stop = stop - last_start

    before_expected = full
    after_expected = full[:start] + new + full[stop:]

    if first_i == last_i:
        r = runs[first_i]
        r.text = (
            r.text[:local_start]
            + new
            + r.text[local_start + len(old):]
        )

    else:
        first = runs[first_i]
        last = runs[last_i]

        prefix = first.text[:local_start]
        suffix = last.text[local_stop:]

        first.text = prefix + new

        for i in touched[1:-1]:
            runs[i].text = ""

        last.text = suffix

    after_actual = "".join(r.text for r in runs)

    if after_actual != after_expected:
        # This should never happen; hard failure is safer.
        raise RuntimeError(
            "POST_REPLACEMENT_TEXT_MISMATCH"
        )

    return True, "APPLIED"


print()
print("=" * 92)
print("STEP23C-QA2B1A: APPLY SAFE LOCKED EDITS")
print("=" * 92)
print()

for p in [QA1B, WORKING, QA2G_LOCK, CONTEXT_MAP]:
    if not p.exists():
        raise SystemExit(f"STOP: missing required file: {p}")

qa1b_sha = sha256(QA1B)

if qa1b_sha != EXPECTED_QA1B_SHA:
    raise SystemExit("STOP: frozen QA1B SHA256 changed.")

working_sha = sha256(WORKING)

if working_sha != qa1b_sha:
    raise SystemExit(
        "STOP: QA2B WORKING is no longer byte-identical to frozen QA1B."
    )

print("✓ Frozen QA1B and clean QA2B working copy verified.")
print()

locks = {
    r["item_id"].strip(): r
    for r in read_csv(QA2G_LOCK)
}

maps = {
    r["item_id"].strip(): r
    for r in read_csv(CONTEXT_MAP)
}

missing = sorted(
    x for x in SAFE_IDS
    if x not in locks or x not in maps
)

if missing:
    raise SystemExit(
        f"STOP: missing safe IDs in lock/map: {missing}"
    )

shutil.copy2(WORKING, OUT_DOCX)

doc = Document(OUT_DOCX)

before_texts = [
    p.text
    for p in doc.paragraphs
]

audit = []
applied_ids = []
skipped_ids = []

for item_id in sorted(
    SAFE_IDS,
    key=lambda x: int(x.split("-")[-1])
):
    lock = locks[item_id]
    mp = maps[item_id]

    old = (lock.get("original_text") or "").strip()
    new = (
        lock.get("approved_text_or_instruction")
        or ""
    ).strip()

    idx_text = (
        mp.get("docx_paragraph_index1")
        or ""
    ).strip()

    if not idx_text.isdigit():
        audit.append({
            "item_id": item_id,
            "status": "SKIPPED",
            "reason": "NO_DOCX_PARAGRAPH_INDEX",
            "paragraph": idx_text,
            "old": old,
            "new": new,
        })
        skipped_ids.append(item_id)
        continue

    idx1 = int(idx_text)

    if not (1 <= idx1 <= len(doc.paragraphs)):
        audit.append({
            "item_id": item_id,
            "status": "SKIPPED",
            "reason": "PARAGRAPH_INDEX_OUT_OF_RANGE",
            "paragraph": idx1,
            "old": old,
            "new": new,
        })
        skipped_ids.append(item_id)
        continue

    p = doc.paragraphs[idx1 - 1]

    ok, reason = replace_once_preserve_runs(
        p,
        old,
        new,
    )

    audit.append({
        "item_id": item_id,
        "status": "APPLIED" if ok else "SKIPPED",
        "reason": reason,
        "paragraph": idx1,
        "old": old,
        "new": new,
    })

    if ok:
        applied_ids.append(item_id)
    else:
        skipped_ids.append(item_id)


doc.save(OUT_DOCX)


# ------------------------------------------------------------
# Structural verification:
# only paragraphs containing successful edits may differ.
# ------------------------------------------------------------

check_doc = Document(OUT_DOCX)

after_texts = [
    p.text
    for p in check_doc.paragraphs
]

if len(after_texts) != len(before_texts):
    raise SystemExit(
        "CRITICAL STOP: paragraph count changed."
    )

changed_paras = {
    i + 1
    for i, (a, b) in enumerate(
        zip(before_texts, after_texts)
    )
    if a != b
}

expected_changed = {
    int(maps[x]["docx_paragraph_index1"])
    for x in applied_ids
}

unexpected = changed_paras - expected_changed

if unexpected:
    raise SystemExit(
        "CRITICAL STOP: unexpected paragraphs changed: "
        + repr(sorted(unexpected))
    )


# ------------------------------------------------------------
# Audit CSV
# ------------------------------------------------------------

with open(
    OUT_CSV,
    "w",
    encoding="utf-8-sig",
    newline=""
) as f:
    fields = [
        "item_id",
        "status",
        "reason",
        "paragraph",
        "old",
        "new",
    ]

    w = csv.DictWriter(
        f,
        fieldnames=fields
    )
    w.writeheader()
    w.writerows(audit)


# ------------------------------------------------------------
# Build remaining queue.
# This includes QA2E items, contextual QA2G items,
# QA2C-005, duplicate/merged locks, and package-level work.
# No automatic action is taken on them here.
# ------------------------------------------------------------

all_map_rows = read_csv(CONTEXT_MAP)

remaining = [
    r for r in all_map_rows
    if r["item_id"].strip() not in applied_ids
]

lines = [
    "STEP23C-QA2B1A REMAINING MANUAL/LOCKED QUEUE",
    "=" * 92,
    "",
    f"Applied safe edits: {len(applied_ids)}",
    f"Safe edits skipped: {len(skipped_ids)}",
    f"Remaining mapped/packaging items: {len(remaining)}",
    "",
]

if skipped_ids:
    lines.extend([
        "SAFE ITEMS SKIPPED BY GUARDRAILS:",
        ", ".join(skipped_ids),
        "",
    ])

for r in remaining:
    lines.extend([
        "-" * 92,
        f"ITEM: {r.get('item_id','')}",
        f"SOURCE LOCK: {r.get('source_lock','')}",
        f"SECTION: {r.get('section','')}",
        f"FROZEN PARAGRAPH: {r.get('paragraph','')}",
        f"DOCX PARAGRAPH: {r.get('docx_paragraph_index1','')}",
        f"DECISION: {r.get('decision','')}",
        f"EDIT GROUP: {r.get('edit_group','')}",
        f"MAP STATUS: {r.get('map_status','')}",
        "",
        "ORIGINAL:",
        r.get("original_text", ""),
        "",
        "APPROVED / INSTRUCTION:",
        r.get("approved", ""),
        "",
        "CURRENT DOCX PARAGRAPH:",
        r.get("docx_paragraph_text", ""),
        "",
    ])

OUT_QUEUE.write_text(
    "\n".join(lines),
    encoding="utf-8"
)


# ------------------------------------------------------------
# Source immutability + checksums
# ------------------------------------------------------------

if sha256(QA1B) != EXPECTED_QA1B_SHA:
    raise SystemExit(
        "CRITICAL STOP: QA1B was modified."
    )

with open(
    OUT_SHA,
    "w",
    encoding="utf-8"
) as f:
    for p in [OUT_DOCX, OUT_CSV, OUT_QUEUE]:
        f.write(
            f"{sha256(p)}  {p.relative_to(ROOT)}\n"
        )


print("=" * 92)
print("STEP23C-QA2B1A COMPLETE")
print("=" * 92)
print()

print("QA1B MODIFIED                 : NO")
print("STEP22A RECOMPUTED            : NO")
print("GLOBAL FIND/REPLACE           : NO")
print()

print("SAFE LOCKED ITEMS REQUESTED   :", len(SAFE_IDS))
print("SAFE EDITS APPLIED            :", len(applied_ids))
print("SAFE EDITS SKIPPED            :", len(skipped_ids))
print()

if applied_ids:
    print(
        "APPLIED IDS                   :",
        ", ".join(applied_ids)
    )

if skipped_ids:
    print(
        "SKIPPED IDS                   :",
        ", ".join(skipped_ids)
    )

print()
print("CHANGED DOCX PARAGRAPHS       :", len(changed_paras))
print()

print("OUTPUT DOCX:")
print(" ", OUT_DOCX.relative_to(ROOT))
print()

print("AUDIT:")
print(" ", OUT_CSV.relative_to(ROOT))
print()

print("REMAINING QUEUE:")
print(" ", OUT_QUEUE.relative_to(ROOT))
print()

if skipped_ids:
    print("STATUS:")
    print(
        "STEP23C_QA2B1A_PARTIAL_SAFE_EDIT__"
        "GUARDRAIL_REVIEW_REQUIRED"
    )
else:
    print("STATUS:")
    print(
        "STEP23C_QA2B1A_SAFE_EDIT_PASS"
    )

print()
print("NEXT:")
print(
    "APPLY CONTEXTUAL + QA2E LOCKED DECISIONS "
    "USING THE REMAINING QUEUE; THEN HANDLE PACKAGE ITEMS."
)
print()

