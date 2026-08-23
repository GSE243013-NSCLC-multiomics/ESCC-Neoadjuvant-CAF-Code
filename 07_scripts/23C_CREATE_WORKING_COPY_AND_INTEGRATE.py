#!/usr/bin/env python3

# ============================================================
# ESCC Neoadjuvant Project
# STEP23C
#
# CREATE NEW WORKING COPY + INTEGRATE FROZEN STEP23B TEXT
#
# IMPORTANT
#
# - NEVER overwrite Step21E source manuscript
# - NO Step22A recomputation
# - NO new scientific analysis
# - NO new biological mechanism claim
# - ONLY integrate text frozen in Step23B
#
# Because the source manuscript does not use standard Heading1
# consistently, section detection uses:
#
#   heading TEXT + optional numeric prefix
#
# Examples supported:
#   Methods
#   2. Methods
#   2. Materials and Methods
#   3 Results
#   4. Discussion
#
# If Methods / Results / Discussion cannot be identified
# unambiguously, STOP BEFORE EDITING.
#
# ============================================================

from pathlib import Path
import csv
import hashlib
import re
import shutil
import subprocess
import sys
import zipfile
import tempfile


PROJECT = Path.cwd()

STEP23_META = (
    PROJECT
    / "00_metadata"
    / "Step23_manuscript_update"
)

STEP23_RESULT = (
    PROJECT
    / "04_results"
    / "manuscript_update"
    / "Step23"
)

STEP23_META.mkdir(
    parents=True,
    exist_ok=True
)

STEP23_RESULT.mkdir(
    parents=True,
    exist_ok=True
)

STEP23B_STATUS = (
    STEP23_META
    / "STEP23B_STATUS.csv"
)

STEP23B_TEXT = (
    STEP23_RESULT
    / "STEP23B_PROPOSED_INSERTION_TEXT.txt"
)

STEP23B_RULES = (
    STEP23_RESULT
    / "STEP23B_FROZEN_REPORTING_RULES.txt"
)

WORKING_COPY = (
    STEP23_RESULT
    / "ESCC_Fibroblast_CORE2_Manuscript_STEP23C_WORKING.docx"
)

CHANGELOG_OUT = (
    STEP23_META
    / "STEP23C_MANUSCRIPT_CHANGELOG.csv"
)

SECTION_OUT = (
    STEP23_META
    / "STEP23C_SECTION_RESOLUTION.csv"
)

QA_OUT = (
    STEP23_META
    / "STEP23C_STRUCTURAL_QA.csv"
)

STATUS_OUT = (
    STEP23_META
    / "STEP23C_STATUS.csv"
)

TEXT_SNAPSHOT_OUT = (
    STEP23_RESULT
    / "STEP23C_INSERTED_TEXT_SNAPSHOT.txt"
)

QA_RENDER_DIR = (
    STEP23_RESULT
    / "STEP23C_QA_RENDER"
)


print()
print("=" * 80)
print("STEP23C: CREATE NEW WORKING COPY AND INTEGRATE")
print("=" * 80)
print()


# ============================================================
# Helpers
# ============================================================

def fail(msg):

    print()
    print("ERROR:")
    print(msg)
    print()

    raise SystemExit(1)


def clean(x):

    return (
        str(x or "")
        .replace("\r", "")
        .strip()
    )


def read_csv(path):

    if not path.exists():

        fail(
            "Missing required file:\n"
            + str(path)
        )

    with open(
        path,
        "r",
        newline="",
        encoding="utf-8-sig"
    ) as f:

        return list(
            csv.DictReader(f)
        )


def write_csv(
    path,
    rows,
    fields
):

    with open(
        path,
        "w",
        newline="",
        encoding="utf-8-sig"
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fields,
            extrasaction="ignore"
        )

        writer.writeheader()
        writer.writerows(rows)


def sha256_file(path):

    h = hashlib.sha256()

    with open(
        path,
        "rb"
    ) as f:

        while True:

            b = f.read(
                1024 * 1024
            )

            if not b:
                break

            h.update(b)

    return h.hexdigest()


def relative(path):

    try:

        return str(
            path.relative_to(
                PROJECT
            )
        )

    except Exception:

        return str(path)


def normalize_ws(text):

    return re.sub(
        r"\s+",
        " ",
        clean(text)
    )


# ============================================================
# 1. Verify python-docx
# ============================================================

print("===== 1. PYTHON-DOCX =====")
print()

try:

    import docx

    from docx import Document

except Exception as e:

    fail(
        "python-docx is not available.\n"
        "Run this command once:\n\n"
        "python3 -m pip install --user python-docx\n\n"
        f"Original error: {e}"
    )


print(
    "python-docx:",
    getattr(
        docx,
        "__version__",
        "available"
    )
)

print()


# ============================================================
# 2. Verify Step23B frozen state
# ============================================================

print("===== 2. VERIFY STEP23B =====")
print()

status23b = read_csv(
    STEP23B_STATUS
)

status_map = {
    clean(
        x.get(
            "item"
        )
    ):
    clean(
        x.get(
            "status"
        )
    )
    for x in status23b
}


if (
    status_map.get(
        "manuscript_modified"
    )
    != "NO"
):

    fail(
        "Step23B does not confirm manuscript remained unchanged."
    )


if (
    status_map.get(
        "Step22A_recomputed"
    )
    != "NO"
):

    fail(
        "Step23B does not confirm Step22A remained frozen."
    )


if (
    status_map.get(
        "evidence_hierarchy_frozen"
    )
    != "YES"
):

    fail(
        "Step23B evidence hierarchy is not frozen."
    )


if (
    status_map.get(
        "reporting_rules_frozen"
    )
    != "YES"
):

    fail(
        "Step23B reporting rules are not frozen."
    )


SOURCE = (
    PROJECT
    / status_map[
        "manuscript_source"
    ]
)

SOURCE_SHA_EXPECTED = status_map[
    "manuscript_source_SHA256"
]


if not SOURCE.exists():

    fail(
        "Step21E source manuscript missing:\n"
        + str(SOURCE)
    )


SOURCE_SHA_BEFORE = sha256_file(
    SOURCE
)


if (
    SOURCE_SHA_BEFORE
    != SOURCE_SHA_EXPECTED
):

    fail(
        "SOURCE_MANUSCRIPT_CHANGED_SINCE_STEP23B\n\n"
        f"Expected: {SOURCE_SHA_EXPECTED}\n"
        f"Observed: {SOURCE_SHA_BEFORE}"
    )


if not STEP23B_TEXT.exists():

    fail(
        "Frozen Step23B proposed text is missing."
    )


if not STEP23B_RULES.exists():

    fail(
        "Frozen Step23B reporting rules are missing."
    )


STEP23B_TEXT_SHA = sha256_file(
    STEP23B_TEXT
)


print(
    "Source manuscript:"
)

print(
    " ",
    relative(
        SOURCE
    )
)

print()

print(
    "Source SHA256 verified:",
    SOURCE_SHA_BEFORE
)

print(
    "Step23B text SHA256    :",
    STEP23B_TEXT_SHA
)

print()


# ============================================================
# 3. Frozen insertion text
#
# These are exactly the paragraphs frozen during Step23B.
# We additionally confirm that each occurs in the Step23B
# proposed-text artifact before allowing insertion.
# ============================================================

print("===== 3. VERIFY FROZEN INSERTION TEXT =====")
print()


ABSTRACT_TEXT = (
    "In an independent neoadjuvant single-cell cohort (OMIX005710), "
    "definitive ESCC-only validation was precluded because the single "
    "patient with esophageal adenosquamous carcinoma could not be "
    "identified from the available patient-level metadata. In a "
    "pre-specified secondary histology-sensitivity analysis spanning all "
    "22 possible EASC assignments, the treatment-associated CORE "
    "direction was consistently opposite to that observed in the "
    "discovery cohorts."
)


METHODS_1 = (
    "External validation was additionally assessed using the publicly "
    "available OMIX005710 neoadjuvant single-cell RNA-sequencing cohort. "
    "Fibroblast pseudobulk profiles were constructed from raw counts after "
    "broad cell-type annotation using a marker framework independent of "
    "the CORE genes. A pre-specified minimum of 20 fibroblast cells per "
    "patient-timepoint was required for technical eligibility; one sample "
    "(P15 before treatment, 18 fibroblasts) therefore remained in raw "
    "provenance but was excluded from the technical reference and paired "
    "analysis. Raw pseudobulk counts were transformed to library-size CPM "
    "and log2(CPM+1), followed by gene-wise z-scoring within the eligible "
    "OMIX tumor reference. CORE2_STRONG and CORE4 were calculated using "
    "the previously frozen formulas without refitting gene weights."
)


METHODS_2 = (
    "The source cohort contained one patient with esophageal "
    "adenosquamous carcinoma (EASC), but the corresponding individual "
    "patient could not be resolved from the publicly available metadata. "
    "Accordingly, definitive ESCC-only primary validation was considered "
    "unavailable. Before examining CORE scores, we therefore pre-specified "
    "a secondary histology-sensitivity framework comprising all 22 "
    "possible assignments of the EASC identity. Each scenario excluded "
    "the corresponding patient's technically eligible tumor observations "
    "from the z-score reference and, when applicable, from the paired "
    "analysis. Paired treatment effects were defined as after-treatment "
    "minus before-treatment score. Robustness to individual paired "
    "patients was assessed by full leave-one-patient-out recomputation, "
    "with the omitted patient removed from both the z-score reference and "
    "the paired effect calculation."
)


RESULTS_1 = (
    "Definitive ESCC-only external validation in OMIX005710 was not "
    "possible because the single EASC case could not be mapped to an "
    "individual patient using the available public metadata. We therefore "
    "evaluated the pre-specified secondary histology-sensitivity analysis "
    "across all 22 possible EASC assignments, without selecting or "
    "promoting any individual scenario."
)


RESULTS_2 = (
    "The treatment-associated change in CORE2_STRONG was negative in all "
    "22 scenarios, both for the mean and median paired change, opposite to "
    "the frozen replication direction from the discovery cohorts. CORE4 "
    "showed the same negative mean direction in all 22 scenarios. Both "
    "CDKN1B and GSN contributed to the negative CORE2_STRONG direction in "
    "all 22 scenarios, whereas the BGN/TIMP1 component was negative in 18 "
    "and positive in 4 scenarios. Independent reconstruction of the "
    "normalization, z-scoring, and score calculations reproduced the "
    "original analysis within floating-point precision."
)


RESULTS_3 = (
    "The result was also insensitive to any single paired patient. Across "
    "105 full leave-one-paired-patient-out recomputations, the mean and "
    "median CORE2_STRONG treatment-associated changes remained negative "
    "in every run, and all remaining paired patients retained negative "
    "changes in every run. Thus, within this pre-specified secondary "
    "framework, OMIX005710 showed robust directional discordance rather "
    "than replication of the treatment-associated CORE response observed "
    "in the discovery cohorts."
)


DISCUSSION_1 = (
    "The OMIX005710 analysis provides an important negative external "
    "finding. Although unresolved patient-level histology prevented a "
    "definitive ESCC-only primary validation, the pre-specified "
    "histology-sensitivity analysis consistently yielded a "
    "treatment-associated CORE direction opposite to that observed in "
    "the discovery cohorts. This directional discordance was robust to "
    "all hypothetical EASC assignments, independent reconstruction of the "
    "scoring pipeline, and paired leave-one-patient-out analyses, arguing "
    "against a simple sign, pairing, reference-set, or single-patient "
    "artifact."
)


DISCUSSION_2 = (
    "These findings indicate that the treatment-associated CORE response "
    "is not uniformly replicated across independent cohorts and may be "
    "context-dependent. However, the present analysis does not establish "
    "the biological cause of the discordance. In particular, differences "
    "in CAF state composition, treatment exposure, sampling context, "
    "clinical response, or technical platform were not tested as "
    "confirmatory explanations and would require separately designated "
    "exploratory analyses."
)


LIMITATIONS_TEXT = (
    "A key limitation of the OMIX005710 analysis was the inability to "
    "identify the single EASC patient at the individual-patient level from "
    "the available public metadata, which prevented construction of a "
    "definitive ESCC-only primary validation set. We therefore relied on a "
    "pre-specified secondary sensitivity analysis across all 22 possible "
    "EASC assignments. In addition, only five paired patients satisfied "
    "the frozen fibroblast technical criteria before hypothetical EASC "
    "exclusion, and scenarios in which a paired patient was assigned as "
    "EASC contained four evaluable pairs. This small paired sample size "
    "limits formal inferential resolution; accordingly, interpretation "
    "focused on direction, patient-level consistency, scenario "
    "robustness, and leave-one-patient-out influence rather than "
    "statistical significance."
)


CONCLUSION_TEXT = (
    "Together, the discovery cohorts support a reproducible fibroblast "
    "CORE framework within the original analytical context, whereas the "
    "independent OMIX005710 sensitivity analysis showed robust directional "
    "discordance, indicating that the treatment-associated CORE response "
    "is not uniformly reproduced across cohorts."
)


frozen_text_normalized = normalize_ws(
    STEP23B_TEXT.read_text(
        encoding="utf-8",
        errors="replace"
    )
)


FROZEN_PARAGRAPHS = {
    "ABSTRACT":
        ABSTRACT_TEXT,

    "METHODS_1":
        METHODS_1,

    "METHODS_2":
        METHODS_2,

    "RESULTS_1":
        RESULTS_1,

    "RESULTS_2":
        RESULTS_2,

    "RESULTS_3":
        RESULTS_3,

    "DISCUSSION_1":
        DISCUSSION_1,

    "DISCUSSION_2":
        DISCUSSION_2,

    "LIMITATIONS":
        LIMITATIONS_TEXT,

    "CONCLUSION":
        CONCLUSION_TEXT
}


for name, text in FROZEN_PARAGRAPHS.items():

    present = (
        normalize_ws(
            text
        )
        in
        frozen_text_normalized
    )

    print(
        f"{name:<16s}:",
        "FROZEN"
        if present
        else "NOT FOUND"
    )

    if not present:

        fail(
            "Frozen Step23B paragraph not found: "
            + name
        )


print()
print(
    "✓ All insertion text matches Step23B frozen draft."
)
print()


# ============================================================
# 4. Load document READ-ONLY and resolve structure
# ============================================================

print("===== 4. RESOLVE MANUSCRIPT STRUCTURE =====")
print()

source_doc = Document(
    SOURCE
)


def paragraph_text(p):

    return normalize_ws(
        p.text
    )


def stripped_heading_text(text):

    t = normalize_ws(
        text
    )

    # Remove common numeric prefixes:
    # 2
    # 2.
    # 2.1
    # 2.1.
    # 2)
    # 2.1)
    t = re.sub(
        r"^\s*\d+(?:\.\d+)*\s*[\.\)]?\s*",
        "",
        t
    )

    return t.strip(
        " :.-"
    )


def canonical_heading(text):

    t = stripped_heading_text(
        text
    ).lower()

    mapping = [
        (
            "Abstract",
            r"abstract"
        ),

        (
            "Introduction",
            r"introduction"
        ),

        (
            "Methods",
            r"(?:materials?\s+and\s+methods|methods|methodology)"
        ),

        (
            "Results",
            r"results"
        ),

        (
            "Discussion",
            r"discussion"
        ),

        (
            "Limitations",
            r"(?:limitations?|study\s+limitations?)"
        ),

        (
            "Conclusion",
            r"(?:conclusions?|concluding\s+remarks)"
        ),

        (
            "References",
            r"references"
        )
    ]

    for canonical, pattern in mapping:

        if re.fullmatch(
            pattern,
            t,
            flags=re.I
        ):

            return canonical

    return None


section_hits = []


for i, p in enumerate(
    source_doc.paragraphs
):

    text = paragraph_text(
        p
    )

    if not text:
        continue

    canonical = canonical_heading(
        text
    )

    if canonical:

        section_hits.append({
            "section":
                canonical,

            "paragraph_index_0based":
                i,

            "paragraph_index_1based":
                i + 1,

            "text":
                text,

            "style":
                (
                    p.style.name
                    if p.style is not None
                    else ""
                )
        })


section_map = {}


for row in section_hits:

    section = row[
        "section"
    ]

    if section not in section_map:

        section_map[
            section
        ] = row


for section in [
    "Abstract",
    "Introduction",
    "Methods",
    "Results",
    "Discussion",
    "Limitations",
    "Conclusion",
    "References"
]:

    row = section_map.get(
        section
    )

    if row:

        print(
            f"{section:<14s}: "
            f"paragraph {row['paragraph_index_1based']:<4d} | "
            f"{row['text']}"
        )

    else:

        print(
            f"{section:<14s}: NOT DETECTED"
        )


print()


# Verify that all mandatory sections exist and are unique.
# The DOCUMENT ORDER is not assumed to be fixed:
# journals may place Methods after Discussion.
# Insertion targets are resolved per-section using the
# next major heading that follows that section in document order.
for required in [
    "Abstract",
    "Methods",
    "Results",
    "Discussion",
    "References"
]:

    if required not in section_map:

        print(
            "\nHeading-like short paragraphs for diagnosis:\n"
        )

        for i, p in enumerate(
            source_doc.paragraphs
        ):

            t = paragraph_text(
                p
            )

            if (
                t
                and
                len(t) <= 100
            ):

                style_name = (
                    p.style.name
                    if p.style is not None
                    else ""
                )

                if (
                    re.match(
                        r"^\d",
                        t
                    )
                    or
                    "heading" in style_name.lower()
                    or
                    t.lower()
                    in {
                        "abstract",
                        "introduction",
                        "methods",
                        "results",
                        "discussion",
                        "references"
                    }
                ):

                    print(
                        f"{i + 1:4d} | "
                        f"{style_name:<20s} | "
                        f"{t}"
                    )

        fail(
            "SAFE_SECTION_RESOLUTION_FAILED: "
            + required
        )


# Verify no duplicate mandatory headings.
for required in [
    "Abstract",
    "Methods",
    "Results",
    "Discussion",
    "References"
]:

    occurrences = [
        row
        for row in section_hits
        if row[
            "section"
        ]
        == required
    ]

    if len(occurrences) != 1:

        fail(
            f"Non-unique {required} heading: "
            f"{len(occurrences)} found."
        )


section_rows = []

for row in section_hits:

    section_rows.append({
        "section":
            row[
                "section"
            ],

        "paragraph_index":
            row[
                "paragraph_index_1based"
            ],

        "heading_text":
            row[
                "text"
            ],

        "style":
            row[
                "style"
            ]
    })


write_csv(
    SECTION_OUT,
    section_rows,
    [
        "section",
        "paragraph_index",
        "heading_text",
        "style"
    ]
)


print(
    "✓ Required manuscript sections resolved safely."
)

print()


# ============================================================
# 5. Source structural fingerprint
# ============================================================

print("===== 5. SOURCE STRUCTURAL FINGERPRINT =====")
print()


def image_relation_count(doc):

    count = 0

    for rel in doc.part.rels.values():

        reltype = str(
            rel.reltype
        )

        if reltype.endswith(
            "/image"
        ):

            count += 1

    return count


SOURCE_STATS = {
    "paragraphs":
        len(
            source_doc.paragraphs
        ),

    "tables":
        len(
            source_doc.tables
        ),

    "sections":
        len(
            source_doc.sections
        ),

    "inline_shapes":
        len(
            source_doc.inline_shapes
        ),

    "image_relationships":
        image_relation_count(
            source_doc
        )
}


for key, value in SOURCE_STATS.items():

    print(
        f"{key:<20s}: {value}"
    )


print()


# ============================================================
# 6. Create NEW working copy
# ============================================================

print("===== 6. CREATE WORKING COPY =====")
print()


if (
    WORKING_COPY.resolve()
    ==
    SOURCE.resolve()
):

    fail(
        "WORKING COPY PATH EQUALS SOURCE PATH"
    )


if WORKING_COPY.exists():

    WORKING_COPY.unlink()


shutil.copy2(
    SOURCE,
    WORKING_COPY
)


if (
    sha256_file(
        WORKING_COPY
    )
    !=
    SOURCE_SHA_BEFORE
):

    fail(
        "INITIAL_WORKING_COPY_NOT_BYTE_IDENTICAL"
    )


print(
    "Working copy created:"
)

print(
    " ",
    relative(
        WORKING_COPY
    )
)

print()

print(
    "✓ Initial copy is byte-identical to Step21E source."
)

print()


# ============================================================
# 7. Reopen WORKING COPY for editing
# ============================================================

doc = Document(
    WORKING_COPY
)


# Re-resolve headings in working copy.
def find_section_paragraph(
    document,
    section_name
):

    hits = []

    for p in document.paragraphs:

        if (
            canonical_heading(
                paragraph_text(
                    p
                )
            )
            ==
            section_name
        ):

            hits.append(
                p
            )

    if len(hits) != 1:

        fail(
            f"Expected exactly one {section_name} heading; "
            f"found {len(hits)}."
        )

    return hits[0]


ABSTRACT_HEADING = find_section_paragraph(
    doc,
    "Abstract"
)

METHODS_HEADING = find_section_paragraph(
    doc,
    "Methods"
)

RESULTS_HEADING = find_section_paragraph(
    doc,
    "Results"
)

DISCUSSION_HEADING = find_section_paragraph(
    doc,
    "Discussion"
)

REFERENCES_HEADING = find_section_paragraph(
    doc,
    "References"
)


def paragraph_index(
    document,
    target
):

    for i, p in enumerate(
        document.paragraphs
    ):

        if p._p is target._p:

            return i

    raise ValueError(
        "Paragraph not found"
    )


def previous_nonempty_paragraph(
    document,
    target
):

    idx = paragraph_index(
        document,
        target
    )

    for j in range(
        idx - 1,
        -1,
        -1
    ):

        p = document.paragraphs[j]

        if normalize_ws(
            p.text
        ):

            return p

    return None


def style_name_of(
    paragraph
):

    if (
        paragraph is None
        or
        paragraph.style is None
    ):

        return None

    return paragraph.style.name


def insert_before(
    target,
    text,
    style_name=None
):

    p = target.insert_paragraph_before()

    if style_name:

        try:

            p.style = style_name

        except Exception:
            pass

    p.add_run(
        text
    )

    return p


def insert_multiple_before(
    target,
    texts,
    style_name=None
):

    inserted = []

    for text in texts:

        p = insert_before(
            target,
            text,
            style_name=style_name
        )

        inserted.append(
            p
        )

    return inserted


# ============================================================
# 7b. Resolve successor headings for dynamic insertion
#
# Each section's insertion target is the NEXT major heading
# that follows it in document order. This is robust to any
# journal-specific ordering (e.g., Methods after Discussion).
# ============================================================


def successor_heading_after(
    document,
    heading
):

    idx = paragraph_index(
        document,
        heading
    )

    for p in document.paragraphs[
        idx + 1:
    ]:

        if canonical_heading(
            paragraph_text(
                p
            )
        ):

            return p

    return None


ABSTRACT_HEADING = find_section_paragraph(
    doc,
    "Abstract"
)

RESULTS_HEADING = find_section_paragraph(
    doc,
    "Results"
)

DISCUSSION_HEADING = find_section_paragraph(
    doc,
    "Discussion"
)

METHODS_HEADING = find_section_paragraph(
    doc,
    "Methods"
)

REFERENCES_HEADING = find_section_paragraph(
    doc,
    "References"
)

INTRODUCTION_HEADING = None

try:

    INTRODUCTION_HEADING = find_section_paragraph(
        doc,
        "Introduction"
    )

except SystemExit:

    raise

except Exception:

    INTRODUCTION_HEADING = None


# ============================================================
# 8. ABSTRACT insertion position
#
# Prefer immediately before Keywords.
# Otherwise immediately before the heading that follows
# the Abstract in document order.
# ============================================================

print("===== 7. INTEGRATE ABSTRACT =====")
print()


abstract_successor = successor_heading_after(
    doc,
    ABSTRACT_HEADING
)


abs_idx = paragraph_index(
    doc,
    ABSTRACT_HEADING
)

abs_boundary_idx = paragraph_index(
    doc,
    abstract_successor
)

keyword_target = None


for p in doc.paragraphs[
    abs_idx + 1:
    abs_boundary_idx
]:

    if re.match(
        r"^\s*keywords?\s*[:：]",
        normalize_ws(
            p.text
        ),
        flags=re.I
    ):

        keyword_target = p
        break


abstract_target = (
    keyword_target
    if keyword_target is not None
    else abstract_successor
)


abstract_style = style_name_of(
    previous_nonempty_paragraph(
        doc,
        abstract_target
    )
)


insert_before(
    abstract_target,
    ABSTRACT_TEXT,
    style_name=abstract_style
)


print(
    "✓ Abstract sentence inserted."
)

print()


# ============================================================
# 9. METHODS insertion
#
# Insert at end of Methods: directly before the heading that
# follows Methods in document order (commonly References).
# ============================================================

print("===== 8. INTEGRATE METHODS =====")
print()


methods_successor = successor_heading_after(
    doc,
    METHODS_HEADING
)

methods_body_style = style_name_of(
    previous_nonempty_paragraph(
        doc,
        methods_successor
    )
)


insert_multiple_before(
    methods_successor,
    [
        METHODS_1,
        METHODS_2
    ],
    style_name=methods_body_style
)


print(
    "✓ Methods paragraphs inserted:",
    2
)

print()


# ============================================================
# 10. RESULTS insertion
#
# Insert at end of Results, before the heading that follows
# Results (commonly Discussion).
# ============================================================

print("===== 9. INTEGRATE RESULTS =====")
print()


results_successor = successor_heading_after(
    doc,
    RESULTS_HEADING
)

results_body_style = style_name_of(
    previous_nonempty_paragraph(
        doc,
        results_successor
    )
)


insert_multiple_before(
    results_successor,
    [
        RESULTS_1,
        RESULTS_2,
        RESULTS_3
    ],
    style_name=results_body_style
)


print(
    "✓ Results paragraphs inserted:",
    3
)

print()


# ============================================================
# 11. DISCUSSION / LIMITATIONS / CONCLUSION
#
# Discussion text is inserted before the heading that follows
# the Discussion heading in document order.
#
# If the source has no dedicated Limitations or Conclusion
# heading, those paragraphs are appended to the Discussion
# (still before Discussion's successor) without inventing
# new major headings.
# ============================================================

print("===== 10. INTEGRATE DISCUSSION / LIMITATIONS / CONCLUSION =====")
print()


# Re-resolve after earlier insertions.
def optional_section(
    document,
    section_name
):

    hits = [
        p
        for p in document.paragraphs
        if (
            canonical_heading(
                paragraph_text(
                    p
                )
            )
            ==
            section_name
        )
    ]

    if len(hits) == 1:

        return hits[0]

    if len(hits) == 0:

        return None

    fail(
        f"Multiple {section_name} headings detected."
    )


LIMITATIONS_HEADING = optional_section(
    doc,
    "Limitations"
)

CONCLUSION_HEADING = optional_section(
    doc,
    "Conclusion"
)


# The heading that follows the Discussion section in doc order.
discussion_successor = successor_heading_after(
    doc,
    DISCUSSION_HEADING
)

discussion_style = style_name_of(
    previous_nonempty_paragraph(
        doc,
        discussion_successor
    )
)


insert_multiple_before(
    discussion_successor,
    [
        DISCUSSION_1,
        DISCUSSION_2
    ],
    style_name=discussion_style
)


if LIMITATIONS_HEADING is not None:

    limitation_successor = successor_heading_after(
        doc,
        LIMITATIONS_HEADING
    )

    limitation_style = style_name_of(
        previous_nonempty_paragraph(
            doc,
            limitation_successor
        )
    )

    insert_before(
        limitation_successor,
        LIMITATIONS_TEXT,
        style_name=limitation_style
    )

    limitation_location = (
        "EXISTING_LIMITATIONS_SECTION"
    )

else:

    # No dedicated Limitations heading:
    # append limitation paragraph to Discussion.
    insert_before(
        discussion_successor,
        LIMITATIONS_TEXT,
        style_name=discussion_style
    )

    limitation_location = (
        "DISCUSSION_END_NO_NEW_HEADING"
    )


if CONCLUSION_HEADING is not None:

    conclusion_successor = successor_heading_after(
        doc,
        CONCLUSION_HEADING
    )

    conclusion_style = style_name_of(
        previous_nonempty_paragraph(
            doc,
            conclusion_successor
        )
    )

    insert_before(
        conclusion_successor,
        CONCLUSION_TEXT,
        style_name=conclusion_style
    )

    conclusion_location = (
        "EXISTING_CONCLUSION_SECTION"
    )

else:

    # Do not invent a new major heading.
    # Place the calibration sentence at the end of Discussion,
    # after the limitation paragraph.
    insert_before(
        discussion_successor,
        CONCLUSION_TEXT,
        style_name=discussion_style
    )

    conclusion_location = (
        "DISCUSSION_END_NO_NEW_HEADING"
    )


print(
    "✓ Discussion paragraphs inserted: 2"
)

print(
    "✓ Limitation placement:",
    limitation_location
)

print(
    "✓ Conclusion placement:",
    conclusion_location
)

print()


# ============================================================
# 12. Save WORKING COPY only
# ============================================================

print("===== 11. SAVE WORKING COPY =====")
print()


doc.save(
    WORKING_COPY
)


if not WORKING_COPY.exists():

    fail(
        "Working DOCX was not written."
    )


SOURCE_SHA_AFTER = sha256_file(
    SOURCE
)

OUTPUT_SHA = sha256_file(
    WORKING_COPY
)


if (
    SOURCE_SHA_AFTER
    != SOURCE_SHA_BEFORE
):

    fail(
        "CRITICAL: STEP21E SOURCE MANUSCRIPT WAS MODIFIED"
    )


if (
    OUTPUT_SHA
    ==
    SOURCE_SHA_BEFORE
):

    fail(
        "Working copy appears unchanged after integration."
    )


print(
    "✓ Working copy saved."
)

print(
    "✓ Step21E source SHA256 unchanged."
)

print()


# ============================================================
# 13. Text verification
# ============================================================

print("===== 12. INSERTION TEXT VERIFICATION =====")
print()


check_doc = Document(
    WORKING_COPY
)

full_text = "\n".join(
    p.text
    for p in check_doc.paragraphs
)

full_norm = normalize_ws(
    full_text
)


insert_counts = {}


for name, text in FROZEN_PARAGRAPHS.items():

    needle = normalize_ws(
        text
    )

    count = full_norm.count(
        needle
    )

    insert_counts[
        name
    ] = count

    print(
        f"{name:<16s}: {count}"
    )

    if count != 1:

        fail(
            f"Expected inserted text exactly once: {name}; "
            f"observed {count}."
        )


print()
print(
    "✓ All frozen insertion paragraphs occur exactly once."
)

print()


# ============================================================
# 14. Structural preservation audit
# ============================================================

print("===== 13. STRUCTURAL PRESERVATION AUDIT =====")
print()


OUTPUT_STATS = {
    "paragraphs":
        len(
            check_doc.paragraphs
        ),

    "tables":
        len(
            check_doc.tables
        ),

    "sections":
        len(
            check_doc.sections
        ),

    "inline_shapes":
        len(
            check_doc.inline_shapes
        ),

    "image_relationships":
        image_relation_count(
            check_doc
        )
}


EXPECTED_NEW_PARAGRAPHS = 10


qa_rows = []


def add_qa(
    check,
    observed,
    expected,
    passed
):

    qa_rows.append({
        "check":
            check,

        "observed":
            str(
                observed
            ),

        "expected":
            str(
                expected
            ),

        "status":
            (
                "PASS"
                if passed
                else "FAIL"
            )
    })


add_qa(
    "source_SHA_unchanged",
    SOURCE_SHA_AFTER,
    SOURCE_SHA_BEFORE,
    (
        SOURCE_SHA_AFTER
        ==
        SOURCE_SHA_BEFORE
    )
)


add_qa(
    "paragraph_increment",
    (
        OUTPUT_STATS[
            "paragraphs"
        ]
        -
        SOURCE_STATS[
            "paragraphs"
        ]
    ),
    EXPECTED_NEW_PARAGRAPHS,
    (
        OUTPUT_STATS[
            "paragraphs"
        ]
        -
        SOURCE_STATS[
            "paragraphs"
        ]
        ==
        EXPECTED_NEW_PARAGRAPHS
    )
)


for key in [
    "tables",
    "sections",
    "inline_shapes",
    "image_relationships"
]:

    add_qa(
        key
        + "_preserved",

        OUTPUT_STATS[
            key
        ],

        SOURCE_STATS[
            key
        ],

        (
            OUTPUT_STATS[
                key
            ]
            ==
            SOURCE_STATS[
                key
            ]
        )
    )


# ============================================================
# 15. Package-level audit
#
# Ensure important Word parts present in source remain present.
# ============================================================

def zip_members(path):

    with zipfile.ZipFile(
        path,
        "r"
    ) as z:

        return set(
            z.namelist()
        )


source_members = zip_members(
    SOURCE
)

output_members = zip_members(
    WORKING_COPY
)


critical_prefixes = [
    "word/media/",
    "word/header",
    "word/footer"
]

critical_exact = [
    "word/styles.xml",
    "word/settings.xml",
    "word/numbering.xml",
    "word/_rels/document.xml.rels"
]


critical_source_parts = set()


for name in source_members:

    if any(
        name.startswith(
            prefix
        )
        for prefix in critical_prefixes
    ):

        critical_source_parts.add(
            name
        )


for name in critical_exact:

    if name in source_members:

        critical_source_parts.add(
            name
        )


for optional in [
    "word/comments.xml",
    "word/footnotes.xml",
    "word/endnotes.xml"
]:

    if optional in source_members:

        critical_source_parts.add(
            optional
        )


missing_critical = sorted(
    critical_source_parts
    -
    output_members
)


add_qa(
    "critical_package_parts_missing",
    len(
        missing_critical
    ),
    0,
    (
        len(
            missing_critical
        )
        ==
        0
    )
)


# ============================================================
# 16. XML object-count audit
#
# Detect accidental loss of common complex Word structures.
# ============================================================

def document_xml_bytes(path):

    with zipfile.ZipFile(
        path,
        "r"
    ) as z:

        return z.read(
            "word/document.xml"
        )


source_xml = document_xml_bytes(
    SOURCE
)

output_xml = document_xml_bytes(
    WORKING_COPY
)


xml_tokens = {
    "field_characters":
        b"w:fldChar",

    "hyperlinks":
        b"w:hyperlink",

    "tracked_insertions":
        b"<w:ins",

    "tracked_deletions":
        b"<w:del",

    "math_objects":
        b"m:oMath"
}


for name, token in xml_tokens.items():

    source_count = source_xml.count(
        token
    )

    output_count = output_xml.count(
        token
    )

    add_qa(
        name
        + "_preserved",

        output_count,

        source_count,

        (
            output_count
            ==
            source_count
        )
    )


write_csv(
    QA_OUT,
    qa_rows,
    [
        "check",
        "observed",
        "expected",
        "status"
    ]
)


qa_failures = [
    x
    for x in qa_rows
    if x[
        "status"
    ]
    != "PASS"
]


for row in qa_rows:

    print(
        f"{row['check']:<36s} "
        f"{row['status']:<5s} "
        f"observed={row['observed']} "
        f"expected={row['expected']}"
    )


print()


if missing_critical:

    print(
        "Missing critical package parts:"
    )

    for x in missing_critical:

        print(
            " ",
            x
        )


if qa_failures:

    fail(
        "STEP23C_STRUCTURAL_QA_FAILED"
    )


print(
    "✓ DOCX structural preservation audit passed."
)

print()


# ============================================================
# 17. Changelog
# ============================================================

print("===== 14. WRITE CHANGELOG =====")
print()


change_rows = [
    {
        "section":
            "Abstract",

        "action":
            "INSERT",

        "paragraphs_added":
            1,

        "content":
            "OMIX005710 primary-unavailable + secondary opposite-direction summary"
    },

    {
        "section":
            "Methods",

        "action":
            "INSERT_AT_SECTION_END",

        "paragraphs_added":
            2,

        "content":
            "OMIX technical processing + frozen histology-sensitivity design"
    },

    {
        "section":
            "Results",

        "action":
            "INSERT_AT_SECTION_END",

        "paragraphs_added":
            3,

        "content":
            "22-scenario non-replication + CORE4/components + 105 LOPO"
    },

    {
        "section":
            "Discussion",

        "action":
            "INSERT_AT_SECTION_END",

        "paragraphs_added":
            2,

        "content":
            "Directional discordance interpretation + mechanism boundary"
    },

    {
        "section":
            "Limitations",

        "action":
            "INSERT",

        "paragraphs_added":
            1,

        "content":
            (
                "EASC identity limitation + paired small-n inferential limitation"
            )
    },

    {
        "section":
            "Conclusion",

        "action":
            "CALIBRATION_INSERT",

        "paragraphs_added":
            1,

        "content":
            "Cross-cohort reproducibility conclusion recalibration"
    }
]


write_csv(
    CHANGELOG_OUT,
    change_rows,
    [
        "section",
        "action",
        "paragraphs_added",
        "content"
    ]
)


# ============================================================
# 18. Snapshot inserted text
# ============================================================

snapshot = f"""STEP23C INSERTED TEXT SNAPSHOT
======================================================================

SOURCE MANUSCRIPT:
{relative(SOURCE)}

SOURCE SHA256:
{SOURCE_SHA_BEFORE}

STEP23B FROZEN TEXT SHA256:
{STEP23B_TEXT_SHA}

WORKING COPY:
{relative(WORKING_COPY)}

WORKING COPY SHA256:
{OUTPUT_SHA}

======================================================================
ABSTRACT
======================================================================

{ABSTRACT_TEXT}

======================================================================
METHODS
======================================================================

{METHODS_1}

{METHODS_2}

======================================================================
RESULTS
======================================================================

{RESULTS_1}

{RESULTS_2}

{RESULTS_3}

======================================================================
DISCUSSION
======================================================================

{DISCUSSION_1}

{DISCUSSION_2}

======================================================================
LIMITATIONS
======================================================================

{LIMITATIONS_TEXT}

======================================================================
CONCLUSION CALIBRATION
======================================================================

{CONCLUSION_TEXT}
"""


TEXT_SNAPSHOT_OUT.write_text(
    snapshot,
    encoding="utf-8"
)


# ============================================================
# 19. Attempt automatic render for visual QA
#
# Non-blocking if LibreOffice is absent.
# Manuscript is NOT final until visual QA is reviewed.
# ============================================================

print("===== 15. ATTEMPT DOCX RENDER =====")
print()


QA_RENDER_DIR.mkdir(
    parents=True,
    exist_ok=True
)


soffice_candidates = [
    shutil.which(
        "libreoffice"
    ),

    shutil.which(
        "soffice"
    ),

    (
        "/Applications/LibreOffice.app/"
        "Contents/MacOS/soffice"
    )
]


soffice = None


for candidate in soffice_candidates:

    if not candidate:
        continue

    candidate_path = Path(
        candidate
    )

    if candidate_path.exists():

        soffice = str(
            candidate_path
        )

        break


render_status = (
    "NOT_ATTEMPTED_LIBREOFFICE_NOT_FOUND"
)

render_pdf = ""


if soffice:

    try:

        profile_dir = (
            QA_RENDER_DIR
            / "lo_profile"
        )

        profile_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        profile_uri = (
            profile_dir
            .resolve()
            .as_uri()
        )

        cmd = [
            soffice,
            (
                "-env:UserInstallation="
                + profile_uri
            ),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(
                QA_RENDER_DIR
            ),
            str(
                WORKING_COPY
            )
        ]

        print(
            "Running:"
        )

        print(
            " ",
            " ".join(
                cmd
            )
        )

        render = subprocess.run(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            timeout=180
        )

        expected_pdf = (
            QA_RENDER_DIR
            /
            (
                WORKING_COPY.stem
                + ".pdf"
            )
        )

        if (
            render.returncode == 0
            and
            expected_pdf.exists()
            and
            expected_pdf.stat().st_size > 0
        ):

            render_status = (
                "PDF_RENDERED"
            )

            render_pdf = relative(
                expected_pdf
            )

            print()
            print(
                "✓ PDF render created:"
            )

            print(
                " ",
                render_pdf
            )

            # Optional PNG conversion with pdftoppm.
            pdftoppm = shutil.which(
                "pdftoppm"
            )

            if pdftoppm:

                png_prefix = (
                    QA_RENDER_DIR
                    / "page"
                )

                png_cmd = [
                    pdftoppm,
                    "-png",
                    "-r",
                    "120",
                    str(
                        expected_pdf
                    ),
                    str(
                        png_prefix
                    )
                ]

                png_render = subprocess.run(
                    png_cmd,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    timeout=180
                )

                png_files = sorted(
                    QA_RENDER_DIR.glob(
                        "page-*.png"
                    )
                )

                if (
                    png_render.returncode == 0
                    and
                    png_files
                ):

                    render_status = (
                        "PDF_AND_PNG_RENDERED"
                    )

                    print(
                        "✓ PNG pages created:",
                        len(
                            png_files
                        )
                    )

        else:

            render_status = (
                "LIBREOFFICE_RENDER_FAILED"
            )

            print(
                render.stdout
            )

    except Exception as e:

        render_status = (
            "LIBREOFFICE_RENDER_EXCEPTION"
        )

        print(
            "Render exception:",
            e
        )

else:

    print(
        "LibreOffice not found; render deferred to visual-QA step."
    )


print()


# ============================================================
# 20. Final status
# ============================================================

status_rows = [
    {
        "item":
            "source_manuscript",

        "status":
            relative(
                SOURCE
            )
    },

    {
        "item":
            "source_SHA256_before",

        "status":
            SOURCE_SHA_BEFORE
    },

    {
        "item":
            "source_SHA256_after",

        "status":
            SOURCE_SHA_AFTER
    },

    {
        "item":
            "source_modified",

        "status":
            "NO"
    },

    {
        "item":
            "working_copy",

        "status":
            relative(
                WORKING_COPY
            )
    },

    {
        "item":
            "working_copy_SHA256",

        "status":
            OUTPUT_SHA
    },

    {
        "item":
            "Step22A_recomputed",

        "status":
            "NO"
    },

    {
        "item":
            "frozen_Step23B_text_used",

        "status":
            "YES"
    },

    {
        "item":
            "new_analysis_performed",

        "status":
            "NO"
    },

    {
        "item":
            "structural_QA",

        "status":
            "PASS"
    },

    {
        "item":
            "render_status",

        "status":
            render_status
    },

    {
        "item":
            "render_pdf",

        "status":
            render_pdf
    },

    {
        "item":
            "manuscript_status",

        "status":
            "WORKING_COPY_NOT_FINAL"
    },

    {
        "item":
            "next_step",

        "status":
            "STEP23C_QA_VISUAL_AND_TEXTUAL_REVIEW"
    }
]


write_csv(
    STATUS_OUT,
    status_rows,
    [
        "item",
        "status"
    ]
)


print("=" * 80)
print("STEP23C COMPLETE")
print("=" * 80)
print()

print(
    "STEP21E SOURCE MODIFIED       : NO"
)

print(
    "STEP22A RECOMPUTED            : NO"
)

print(
    "NEW ANALYSIS PERFORMED        : NO"
)

print()

print(
    "WORKING COPY:"
)

print(
    " ",
    relative(
        WORKING_COPY
    )
)

print()

print(
    "FROZEN PARAGRAPHS INSERTED    : 10 / 10"
)

print(
    "STRUCTURAL QA FAILURES        : 0"
)

print()

print(
    "TABLE COUNT PRESERVED         :",
    SOURCE_STATS[
        "tables"
    ],
    "->",
    OUTPUT_STATS[
        "tables"
    ]
)

print(
    "SECTION COUNT PRESERVED       :",
    SOURCE_STATS[
        "sections"
    ],
    "->",
    OUTPUT_STATS[
        "sections"
    ]
)

print(
    "INLINE SHAPES PRESERVED       :",
    SOURCE_STATS[
        "inline_shapes"
    ],
    "->",
    OUTPUT_STATS[
        "inline_shapes"
    ]
)

print(
    "IMAGE RELATIONSHIPS PRESERVED :",
    SOURCE_STATS[
        "image_relationships"
    ],
    "->",
    OUTPUT_STATS[
        "image_relationships"
    ]
)

print()

print(
    "LIMITATIONS PLACEMENT         :",
    limitation_location
)

print(
    "CONCLUSION PLACEMENT          :",
    conclusion_location
)

print()

print(
    "RENDER STATUS                 :",
    render_status
)

if render_pdf:

    print(
        "RENDERED PDF:"
    )

    print(
        " ",
        render_pdf
    )


print()

print(
    "MANUSCRIPT STATUS:"
)

print(
    "WORKING_COPY_NOT_FINAL"
)

print()

print(
    "NEXT:"
)

print(
    "STEP23C_QA_VISUAL_AND_TEXTUAL_REVIEW"
)

print()

print(
    "DO NOT DELETE OR OVERWRITE THE STEP21E SOURCE."
)

print()
